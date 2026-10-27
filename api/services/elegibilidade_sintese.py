"""Situação de cada demanda quanto a restrições e riscos, para o relatório e o pacote da Elegibilidade Territorial.

A demanda é o objeto da análise; as camadas de restrição e de risco apenas a caracterizam. Cada demanda tem
restrição, tem risco, tem ambos (uma situação, não um conceito) ou está livre de restrições e riscos.
As regras espelham as do painel interativo (cinco-resultados_elegibilidade_v2-dados.js).
"""
from __future__ import annotations

import csv
from decimal import Decimal
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.platypus import Flowable

from api.services import extracao_atributos_relatorios as estilo

CATEGORIAS = ('restricao', 'risco')
ROTULOS = {'restricao': 'Restrição', 'risco': 'Risco'}
SITUACOES = {'ambas': 'Com restrição e risco', 'restricao': 'Com restrição', 'risco': 'Com risco',
             'livre': 'Livre de restrições e riscos'}
ORDEM_SITUACOES = ('ambas', 'restricao', 'risco', 'livre')
CORES_SITUACAO = {'ambas': '#7a4a1d', 'restricao': '#b42335', 'risco': '#c18b00', 'livre': '#2aa358'}
PREFIXOS_FORMULA = ('=', '+', '-', '@', '\t', '\r')
CONCEITO = ('A demanda é o objeto da análise; as bases de restrição e de risco servem para caracterizá-la. '
            'Para cada demanda, este relatório informa se ela tem restrição, tem risco, tem ambos ou está livre de '
            'restrições e riscos ambientais e territoriais, conforme as camadas consultadas.')
NOTA_LEITURA = ('Livre de restrições e riscos significa que não foi encontrada sobreposição com as camadas consultadas. '
                'O percentual usa a área total da demanda como base; áreas atingidas por camadas diferentes podem coincidir '
                'e não são somadas. “Apenas toca o limite” indica contato sem área atingida. '
                'O relatório informa a situação de cada demanda; não a classifica como elegível ou inelegível.')
COLUNAS_RESUMO = ['Conjunto de demandas', 'Objeto de demanda', 'Camadas de risco', 'Área de sobreposição (ha)',
                  'Porcentagem sobreposta', 'Índice de risco ambiental e territorial']
COLUNAS_VINCULO = ['Execução', 'Chave da demanda', 'Identificador da demanda']
NOTA_CONFERENCIA = (
    'Use execução + chave posicional da demanda para distinguir demandas, inclusive homônimas. '
    'No detalhe, acrescente o ID da camada e o ID da feição de origem. Nomes são de apresentação. '
    'As tabelas preservam a precisão numérica disponível; a formatação de leitura não modifica os valores. '
    'O XLSX apresenta percentuais como porcentagens; o CSV registra números em 0–100, não frações. '
    'Célula vazia de área indica ausência de interseção, não zero. Contato sem área é uma ocorrência '
    'com área zero e pode caracterizar restrição ou risco. As contagens de ocorrências são de feições, '
    'não necessariamente de camadas distintas. Áreas por camada não são parcelas exclusivas e não devem '
    'ser somadas para recalcular o índice do resumo, que utiliza a união das áreas de risco. '
    'Editar uma cópia da planilha ou CSV não atualiza o SICARD e não constitui reimportação. '
    'Confira a fonte e, se necessário, corrija-a e gere nova execução, preservando o pacote original.'
)
NOTA_INDICE = ('Índice de risco ambiental e territorial = 1 − (área da demanda sobreposta por camadas de risco ÷ área total da demanda). '
               'Varia de 0 a 1: quanto mais próximo de 1, maior a elegibilidade da demanda para execução; quanto mais próximo de 0, '
               'menor. A área sobreposta é a união das áreas atingidas por todas as camadas de risco, sem contar duas vezes o que '
               'coincide. Demandas com restrição não constam desta tabela.')


def indice_risco_ambiental(area_m2: float, total_m2: float) -> float:
    """Índice de 0 a 1: 1 sem sobreposição com camadas de risco; 0 com a demanda inteira sobreposta."""
    if not total_m2 or total_m2 <= 0:
        return 1.0
    return max(0.0, min(1.0, 1 - area_m2 / total_m2))


def caracterizar(entrada: dict) -> list[dict]:
    """Uma ficha por demanda, com o que foi encontrado em cada categoria."""
    camadas = {c['id']: c['nome'] for c in entrada['campos']}
    fichas = []
    for linha in entrada['linhas']:
        cats = {c: [] for c in CATEGORIAS}
        for item in linha['sobreposicoes']:
            if item['papel'] in cats:
                cats[item['papel']].append({**item, 'camada': camadas[item['base_id']]})
        restricao, risco = bool(cats['restricao']), bool(cats['risco'])
        fichas.append({'linha': linha, 'nome': linha['nome'] or f'Demanda {linha["id"]}', 'cats': cats,
                       'situacao': 'ambas' if restricao and risco else 'restricao' if restricao
                       else 'risco' if risco else 'livre'})
    return fichas


def ordenar_por_situacao(fichas: list[dict]) -> list[dict]:
    return sorted(fichas, key=lambda f: (ORDEM_SITUACOES.index(f['situacao']), f['nome'].casefold()))


def indicadores(fichas: list[dict]) -> dict:
    return {'total': len(fichas), **{s: sum(f['situacao'] == s for f in fichas) for s in ORDEM_SITUACOES}}


def nomes_encontrados(itens: list[dict], limite: int | None = None) -> str:
    nomes = list(dict.fromkeys(i['nome'] for i in itens))
    if limite and len(nomes) > limite:
        return '; '.join(nomes[:limite]) + f' e mais {len(nomes) - limite}'
    return '; '.join(nomes)


def descrever(ficha: dict, limite: int | None = None) -> str:
    """O que a demanda tem, em linguagem corrente."""
    if ficha['situacao'] == 'livre':
        return 'Livre de restrições e riscos nas camadas consultadas.'
    plural = {'restricao': ('restrição', 'restrições'), 'risco': ('risco', 'riscos')}
    frases = []
    for c in CATEGORIAS:
        itens = ficha['cats'][c]
        if itens:
            tem = f'Tem {plural[c][0]}' if len(itens) == 1 else f'Tem {len(itens)} {plural[c][1]}'
            frases.append(f'{tem}: {nomes_encontrados(itens, limite)}.')
    return ' '.join(frases)


def registros_detalhados(modelo: dict) -> list[dict]:
    """Formato longo: uma linha por demanda × item encontrado (ou uma linha “livre”)."""
    registros = []
    for entrada in modelo['entradas']:
        for ficha in caracterizar(entrada):
            linha = ficha['linha']
            base = {'entrada': entrada['nome'], 'execucao': entrada['id'], 'chave': linha['chave'],
                    'id': linha['id'], 'demanda': ficha['nome'],
                    'situacao': SITUACOES[ficha['situacao']], 'area_total_m2': linha['area_total_m2']}
            if ficha['situacao'] == 'livre':
                registros.append({**base, 'tipo': None, 'observacao': 'Nenhuma restrição ou risco encontrado nas camadas consultadas'})
                continue
            itens = [i for c in CATEGORIAS for i in ficha['cats'][c]]
            for item in sorted(itens, key=lambda s: -s['area_m2']):
                registros.append({**base, 'tipo': ROTULOS[item['papel']], 'camada': item['camada'], 'item': item['nome'],
                                  'base_id': item['base_id'], 'feicao_id': item['id'],
                                  'area_m2': item['area_m2'], 'percentual': item['percentual'],
                                  'observacao': 'Apenas toca o limite da demanda, sem área atingida' if item['area_m2'] == 0 else ''})
    return registros


def resumo_elegibilidade(entrada: dict) -> list[dict]:
    """Uma linha por demanda sem restrição (com risco ou livre), da mais à menos exposta ao risco."""
    linhas = []
    for posicao, ficha in enumerate(caracterizar(entrada)):
        if ficha['cats']['restricao']:
            continue
        linha, itens = ficha['linha'], sorted(ficha['cats']['risco'], key=lambda i: -i['area_m2'])
        total = linha['area_total_m2']
        area = linha.get('risco_area_m2')
        if area is None:
            area = sum(i['area_m2'] for i in itens)
        area = min(area, total)
        linhas.append({'id': linha['id'], 'chave': str(posicao), 'demanda': ficha['nome'],
                       'camadas': list(dict.fromkeys(i['camada'] for i in itens)),
                       'area_m2': area, 'percentual': area / total * 100 if total else 0.0,
                       'indice': indice_risco_ambiental(area, total)})
    return sorted(linhas, key=lambda r: (r['indice'], r['demanda'].casefold()))


def linhas_resumo(modelo: dict) -> list[dict]:
    return [{**r, 'conjunto': entrada['nome']} for entrada in modelo['entradas'] for r in resumo_elegibilidade(entrada)]


def linhas_resumo_exportacao(modelo: dict) -> list[list]:
    linhas = []
    for entrada in modelo['entradas']:
        for resumo in resumo_elegibilidade(entrada):
            linhas.append([entrada['nome'], resumo['demanda'], '; '.join(resumo['camadas']) or '—',
                           resumo['area_m2'] / 1e4, resumo['percentual'], resumo['indice'],
                           entrada['id'], resumo['chave'], resumo['id']])
    return linhas


def resumo_para_registro(modelo: dict) -> dict:
    """Formato gravado no relatório da Fase 1 da hierarquização (campos com os nomes das colunas da tabela)."""
    return {'titulo': 'Resumo da Elegibilidade', 'nota': NOTA_INDICE, 'colunas': COLUNAS_RESUMO,
            'linhas': [{'conjunto_de_demandas': r['conjunto'], 'objeto_de_demanda': r['demanda'],
                        'camadas_de_risco': r['camadas'], 'area_de_sobreposicao_ha': r['area_m2'] / 1e4,
                        'porcentagem_sobreposta': r['percentual'],
                        'indice_de_risco_ambiental_e_territorial': r['indice']} for r in linhas_resumo(modelo)]}


# ------------------------------------------------------------------ planilhas

def _folha(ws, titulo, colunas, linhas, larguras):
    ws.title = titulo
    ws.append(colunas)
    for celula in ws[1]:
        celula.fill, celula.font = PatternFill('solid', start_color='003B5A'), Font(bold=True, color='FFFFFF')
        celula.alignment = Alignment(wrap_text=True, vertical='center')
    for linha in linhas:
        ws.append(linha)
    for linha in ws.iter_rows(min_row=2):
        for celula in linha:
            if isinstance(celula.value, str):
                celula.data_type = 's'
    for indice, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(indice)].width = largura
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def _folha_resumo_elegibilidade(livro, ws, modelo: dict):
    linhas = linhas_resumo_exportacao(modelo)
    _folha(ws, 'Resumo da Elegibilidade', COLUNAS_RESUMO + COLUNAS_VINCULO,
           [[*r[:4], r[4] / 100, *r[5:]] for r in linhas], [30, 38, 54, 22, 22, 28, 28, 20, 26])
    for linha in ws.iter_rows(min_row=2):
        linha[3].number_format = '0.###############'
        linha[4].number_format = '0.###############%'
        linha[5].number_format = '0.###############'


def _folha_como_ler(livro, modelo: dict):
    ws = livro.create_sheet('Como Ler')
    textos = [
        f'Entradas concluídas: {modelo["concluidas"]} de {modelo["previstas"]}. '
        + ('PACOTE PARCIAL: entradas não concluídas não recebem resultado.'
           if modelo['concluidas'] != modelo['previstas'] else 'Todas as entradas previstas foram concluídas.'),
        CONCEITO, NOTA_LEITURA, NOTA_INDICE, NOTA_CONFERENCIA, modelo['metodologia'],
        'Resumo vazio: nenhuma demanda concluída atende à condição de ausência de restrição. '
        'Isso não significa lote sem demandas nem falha de processamento. Consulte a Situação das Demandas.',
    ]
    for texto in textos:
        ws.append([texto])
        ws.row_dimensions[ws.max_row].height = 110
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 120


def xlsx_resumo(modelo: dict) -> bytes:
    """Planilha só com o Resumo da Elegibilidade, a que fica registrada na hierarquização."""
    livro = Workbook()
    _folha_resumo_elegibilidade(livro, livro.active, modelo)
    _folha_como_ler(livro, modelo)
    memoria = BytesIO()
    livro.save(memoria)
    return memoria.getvalue()


def xlsx(modelo: dict) -> bytes:
    livro = Workbook()
    folha = _folha

    resumo = []
    for entrada in modelo['entradas']:
        for f in ordenar_por_situacao(caracterizar(entrada)):
            linha = f['linha']
            resumo.append([entrada['nome'], linha['id'], f['nome'], SITUACOES[f['situacao']],
                           nomes_encontrados(f['cats']['restricao']) or '—', nomes_encontrados(f['cats']['risco']) or '—',
                           len(f['cats']['restricao']), len(f['cats']['risco']),
                           linha['area_total_m2'] / 1e6, linha['area_total_m2'] / 1e4,
                           entrada['id'], linha['chave']])
    ws = livro.active
    folha(ws, 'Situação das Demandas',
          ['Conjunto de demandas', 'Identificador', 'Demanda', 'Situação', 'Restrições encontradas', 'Riscos encontrados',
           'Nº de restrições', 'Nº de riscos', 'Área da demanda (km²)', 'Área da demanda (ha)',
           'Execução', 'Chave da demanda'],
          resumo, [30, 16, 38, 28, 44, 44, 14, 12, 18, 18, 28, 20])
    for linha in ws.iter_rows(min_row=2):
        linha[8].number_format = '#,##0.00####'
        linha[9].number_format = '#,##0.0000'

    _folha_resumo_elegibilidade(livro, livro.create_sheet(), modelo)

    detalhe = []
    for r in registros_detalhados(modelo):
        livre = r['tipo'] is None
        detalhe.append([r['entrada'], r['id'], r['demanda'], r['situacao'], r['tipo'] or '—',
                        '—' if livre else r['item'], '—' if livre else r['camada'],
                        None if livre else r['area_m2'] / 1e6, None if livre else r['area_m2'] / 1e4,
                        None if livre else r['percentual'] / 100, r['observacao'], r['execucao'], r['chave'],
                        None if livre else r['base_id'], None if livre else r['feicao_id'],
                        None if livre else r['area_m2']])
    ws = livro.create_sheet()
    folha(ws, 'Restrições e Riscos',
          ['Conjunto de demandas', 'Identificador', 'Demanda', 'Situação da demanda', 'Tipo', 'Restrição ou risco encontrado',
           'Camada consultada', 'Área atingida (km²)', 'Área atingida (ha)', '% da área da demanda', 'Observação',
           'Execução', 'Chave da demanda', 'ID da camada', 'ID da feição', 'Área atingida (m²)'],
          detalhe, [30, 16, 38, 28, 12, 40, 36, 18, 18, 20, 44, 28, 20, 28, 28, 24])
    for linha in ws.iter_rows(min_row=2):
        linha[7].number_format = '0.###############'
        linha[8].number_format = '0.###############'
        linha[9].number_format = '0.###############%'
        linha[15].number_format = '0.###############'

    ws = livro.create_sheet()
    folha(ws, 'Camadas Consultadas', ['Papel', 'Nome de apresentação', 'ID da fonte'],
          [[{'entrada': 'Conjunto de demandas', 'restricao': 'Restrição', 'risco': 'Risco'}[f['papel']], f['nome'], f['id']]
           for f in modelo['fontes'] if f['tipo'] == 'camada'], [24, 60, 36])

    _folha_como_ler(livro, modelo)
    memoria = BytesIO()
    livro.save(memoria)
    return memoria.getvalue()


def _celula_csv(valor) -> str:
    if valor is None:
        return ''
    if isinstance(valor, float):
        numero = format(Decimal(str(valor)), 'f')
        return (numero.rstrip('0').rstrip('.') if '.' in numero else numero).replace('.', ',')
    texto = str(valor)
    return "'" + texto if texto.startswith(PREFIXOS_FORMULA) else texto


def csv_detalhe(modelo: dict) -> bytes:
    saida = StringIO()
    escritor = csv.writer(saida, delimiter=';', lineterminator='\r\n')
    escritor.writerow(['Conjunto de demandas', 'Identificador', 'Demanda', 'Situação da demanda', 'Área da demanda (m²)',
                       'Tipo', 'Restrição ou risco encontrado', 'Camada consultada', 'Área atingida (m²)',
                       '% da área da demanda', 'Observação', 'Execução', 'Chave da demanda', 'ID da camada', 'ID da feição'])
    for r in registros_detalhados(modelo):
        livre = r['tipo'] is None
        escritor.writerow([_celula_csv(v) for v in (
            r['entrada'], r['id'], r['demanda'], r['situacao'], r['area_total_m2'], r['tipo'] or '',
            '' if livre else r['item'], '' if livre else r['camada'], None if livre else r['area_m2'],
            None if livre else r['percentual'], r['observacao'], r['execucao'], r['chave'],
            None if livre else r['base_id'], None if livre else r['feicao_id'])])
    return saida.getvalue().encode('utf-8-sig')


def csv_resumo(modelo: dict) -> bytes:
    saida = StringIO()
    escritor = csv.writer(saida, delimiter=';', lineterminator='\r\n')
    escritor.writerow(COLUNAS_RESUMO + COLUNAS_VINCULO)
    for linha in linhas_resumo_exportacao(modelo):
        escritor.writerow([_celula_csv(v) for v in linha])
    return saida.getvalue().encode('utf-8-sig')


def leia_me(modelo: dict, relatorio: str, arquivos: list[str]) -> bytes:
    ind = indicadores([f for e in modelo['entradas'] for f in caracterizar(e)])
    linhas = [
        'RELATÓRIO E CAMADAS — ELEGIBILIDADE TERRITORIAL', '',
        CONCEITO, '',
        f'Demandas analisadas: {ind["total"]}. Com restrição: {ind["restricao"]}. Com risco: {ind["risco"]}. '
        f'Com restrição e risco: {ind["ambas"]}. Livres de restrições e riscos: {ind["livre"]}.',
        f'Conjuntos de demandas concluídos: {modelo["concluidas"]} de {modelo["previstas"]}.',
        'PACOTE PARCIAL: entradas não concluídas não recebem resultado.'
        if modelo['concluidas'] != modelo['previstas'] else 'Todas as entradas previstas foram concluídas.', '',
        'CONTEÚDO DO PACOTE',
        f'- {relatorio}: identificação das execuções, fontes, guia de leitura, metodologia e resultados. '
        'Situação de cada demanda, mapa, resumo e detalhe por demanda; '
        'as matrizes de área por camada ficam no anexo.',
        '- Tabelas/Situação das Demandas.xlsx: situação de cada demanda, restrições e riscos encontrados e camadas consultadas, com filtros.',
        '- Tabelas/Situação das Demandas.csv: formato longo correspondente à aba Restrições e Riscos, '
        'com uma linha por ocorrência ou uma linha sem ocorrência por demanda livre; não é uma linha por demanda.',
        '- Tabelas/Resumo da Elegibilidade.xlsx e .csv: uma linha por demanda sem restrição, com camadas de risco, área e porcentagem '
        'sobrepostas e o índice de risco ambiental e territorial (0 a 1; mais próximo de 1, maior a elegibilidade).',
        '- Camadas/NN/*.gpkg: cópia integral de cada conjunto de demandas, com áreas e sobreposições gravadas na execução (para uso em SIG).',
        '- Metadados da Análise.json: identificação, fontes, nomes revisados e assinatura (SHA-256) dos arquivos.', '',
        'COMO LER', NOTA_LEITURA, NOTA_INDICE, NOTA_CONFERENCIA,
        'CSV: UTF-8 com BOM, separador ponto e vírgula, decimal vírgula. '
        'As áreas do detalhe CSV são em m²; as áreas do resumo são em hectares. '
        'A aba Restrições e Riscos inclui m², km² e hectares. Os percentuais do CSV são em 0–100. '
        'Textos iniciados por caracteres de fórmula recebem apóstrofo de proteção no CSV; '
        'no Excel são células de texto, não fórmulas.',
        'O índice de risco ambiental e territorial descreve a exposição espacial da demanda; '
        'não constitui critério multicritério nem altera pesos AHP, ranking ou homologação. '
        'Se o Resumo da Elegibilidade estiver vazio, consulte a situação de todas as demandas: '
        'demandas com restrição não entram nesse resumo.', '',
        'RASTREABILIDADE E PRÓXIMOS PASSOS',
        'O PDF identifica cada execução e sua data de conclusão. Os metadados contêm as assinaturas SHA-256 '
        'das camadas, PDF, tabelas e guia para conferência. Entradas não concluídas não recebem resultado; '
        'ausência de resultado não significa ausência de restrição ou risco. '
        'Reabra o GeoPackage de saída na seção 5 para consultar novamente, sem repetir o cálculo. '
        'Confira as ocorrências e as fontes antes de avançar à Fase 2 pela navegação da página.', '',
        'ARQUIVOS DE CAMADAS', *[f'- {nome}' for nome in arquivos],
    ]
    return ('\r\n'.join(linhas) + '\r\n').encode('utf-8-sig')


# ------------------------------------------------------------------ PDF

def kpis(e: dict, fichas: list[dict]):
    i = indicadores(fichas)
    return estilo._kpis(e, [(str(i['total']), 'Demandas analisadas'), (str(i['livre']), 'Livres de restrições e riscos'),
                            (str(i['restricao']), 'Com restrição'), (str(i['risco']), 'Com risco'),
                            (str(i['ambas']), 'Com restrição e risco')])


class GraficoSituacoes(Flowable):
    """Barras horizontais com a quantidade de demandas em cada situação."""
    LINHA = 22

    def __init__(self, fichas):
        super().__init__()
        self.contagem = indicadores(fichas)
        self.width = estilo.LARGURA
        self.height = self.LINHA * len(ORDEM_SITUACOES) + 6

    def wrap(self, largura, altura):
        return self.width, self.height

    def draw(self):
        c = self.canv
        x0, x1 = 150, self.width - 70
        total = self.contagem['total'] or 1
        c.saveState()
        for indice, situacao in enumerate(ORDEM_SITUACOES):
            y = self.height - (indice + 1) * self.LINHA
            valor = self.contagem[situacao]
            c.setFillColor(estilo.AZUL)
            c.setFont('Helvetica', 8)
            c.drawString(0, y + 6, SITUACOES[situacao])
            c.setFillColor(colors.HexColor('#e6edf2'))
            c.rect(x0, y + 4, x1 - x0, 10, stroke=0, fill=1)
            if valor:
                c.setFillColor(colors.HexColor(CORES_SITUACAO[situacao]))
                c.rect(x0, y + 4, max(1.5, (x1 - x0) * valor / total), 10, stroke=0, fill=1)
            c.setFillColor(estilo.TEXTO_SUAVE)
            c.setFont('Helvetica', 7.5)
            c.drawString(x1 + 6, y + 6, f'{valor} · {estilo.percentual(valor / total * 100)}')
        c.restoreState()


def tabela_resumo_elegibilidade(e: dict, entrada: dict) -> object:
    linhas = [['Objeto de demanda', 'Camadas de risco', 'Área de sobreposição', 'Porcentagem sobreposta', 'Índice de risco ambiental e territorial']]
    for r in resumo_elegibilidade(entrada):
        linhas.append([r['demanda'], '; '.join(r['camadas']) or '—', f'{estilo.numero(r["area_m2"] / 1e4, 4)} ha',
                       estilo.percentual(r['percentual']), f'{r["indice"]:.4f}'.replace('.', ',')])
    tabela = estilo._tabela(e, linhas, [estilo.LARGURA * p for p in (.27, .27, .16, .14, .16)], direita=(2, 3, 4))
    tabela.splitInRow = 1
    return tabela


def tabela_resumo(e: dict, fichas: list[dict], area_km2) -> object:
    linhas = [['Demanda', 'Situação', 'O que foi encontrado', 'Área da demanda']]
    for f in ordenar_por_situacao(fichas):
        linhas.append([f['nome'], SITUACOES[f['situacao']], descrever(f, 3), area_km2(e, f['linha']['area_total_m2'])])
    tabela = estilo._tabela(e, linhas, [estilo.LARGURA * p for p in (.26, .17, .41, .16)])
    tabela.splitInRow = 1
    return tabela
