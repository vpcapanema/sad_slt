"""Pacote do modo enriquecimento (etapas 6 e 7 do fluxo configurável).

Conteúdo do .zip:

* GeoPackage com uma camada por tipo de geometria da entrada (pontos, linhas,
  polígonos) e a camada de entrada, com os apelidos gravados como nome alternativo
  dos campos;
* uma tabela CSV por camada;
* planilha XLSX com uma aba por camada e a aba ``dicionario_campos``;
* dicionário de campos em CSV (campo, apelido, tema, base, campo de origem, regra);
* ``configuracao.json`` com entrada, bases, regras e procedência, para reproduzir
  o processamento.
"""
from __future__ import annotations

import csv
import io
import json
import tempfile
import zipfile
from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from api.services import extracao_atributos_exportacao as exportacao
from api.services.ciclo_vida_arquivos import apelido
from api.services.extracao_atributos_pacote import _para_gpkg

COLUNAS_DICIONARIO = ['camada', 'campo', 'apelido', 'tema', 'base', 'campo_origem', 'regra']


def nomes(nome_saida: str, camadas, finalidades=None) -> dict[str, tuple[str, str]]:
    base = apelido(nome_saida, 60) or 'extracao'
    arquivos = {'gpkg': (f'{base}.gpkg', 'GeoPackage com uma camada por tipo de geometria, as finalidades e a entrada')}
    for camada in camadas:
        arquivos[f'csv_{camada}'] = (f'{base}_{camada}.csv', f'Tabela de atributos da camada {camada}')
    for chave, item in (finalidades or {}).items():
        for camada in item['camadas']:
            descricao = f"Finalidade {item['nome']}: camada {camada}"
            arquivos[f'csv_{chave}_{camada}'] = (f'{base}_{chave}_{camada}.csv', descricao)
    arquivos['xlsx'] = (f'{base}_tabelas.xlsx', 'Planilha com uma aba por camada e o dicionário de campos')
    arquivos['csv_dicionario'] = (f'{base}_dicionario_campos.csv', 'Dicionário de campos')
    arquivos['configuracao'] = (f'{base}_configuracao.json', 'Configuração usada no processamento')
    arquivos['validacao'] = (f'{base}_validacao.json', 'Conferência do resultado')
    return arquivos


def _texto_xlsx(valor):
    valor = exportacao.celula_xlsx(valor)
    # Caractere de controle (ex.: \x0b vindo de texto colado) é recusado pelo Excel.
    return ILLEGAL_CHARACTERS_RE.sub('', valor) if isinstance(valor, str) else valor


def escrever_gpkg(camadas: dict, entrada, dicionario: list[dict], path: Path, finalidades=None, preservar_geometrias=False, incluir_entrada=True) -> None:
    from osgeo import gdal, ogr
    # nome da camada no GeoPackage -> (tabela, camada de saída de onde vêm os apelidos)
    todas = {nome: (frame, nome) for nome, frame in camadas.items()}
    for chave, item in (finalidades or {}).items():
        for nome, frame in item['camadas'].items():
            todas[f'{chave}_{nome}'] = (frame, nome)
    for nome, (frame, _) in {**todas, **({'entrada': (entrada, None)} if incluir_entrada else {})}.items():
        _para_gpkg(frame.to_crs(4674)).to_file(path, driver='GPKG', layer=nome, engine='pyogrio', index=False,
                                                promote_to_multi=not preservar_geometrias,
                                                **({'geometry_type': 'Unknown'} if preservar_geometrias else {}))
    gdal.UseExceptions()
    fonte = ogr.Open(str(path), 1)
    try:
        for nome, (_, origem) in todas.items():
            camada = fonte.GetLayerByName(nome)
            definicao = camada.GetLayerDefn()
            apelidos_usados = set()
            for item in dicionario:
                if item['camada'] != origem or not item.get('apelido'):
                    continue
                indice = definicao.GetFieldIndex(item['campo'])
                if indice < 0:
                    continue
                campo = ogr.FieldDefn(item['campo'], definicao.GetFieldDefn(indice).GetType())
                apelido_campo = str(item['apelido'])[:255]
                raiz_apelido, numero = apelido_campo, 2
                while apelido_campo.casefold() in apelidos_usados:
                    sufixo = f' ({numero})'
                    apelido_campo = raiz_apelido[:255-len(sufixo)] + sufixo
                    numero += 1
                apelidos_usados.add(apelido_campo.casefold())
                campo.SetAlternativeName(apelido_campo)
                camada.AlterFieldDefn(indice, campo, ogr.ALTER_ALTERNATIVE_NAME_FLAG)
    finally:
        fonte = None


def escrever_dicionario_csv(dicionario: list[dict], path: Path) -> None:
    with open(path, 'w', newline='', encoding='utf-8-sig') as stream:
        escritor = csv.writer(stream, delimiter=';')
        escritor.writerow(COLUNAS_DICIONARIO)
        for item in dicionario:
            escritor.writerow([exportacao.safe(item.get(c)) if item.get(c) is not None else '' for c in COLUNAS_DICIONARIO])


def escrever_xlsx(camadas: dict, dicionario: list[dict], path: Path) -> None:
    livro = Workbook()
    livro.remove(livro.active)
    apelidos = {(d['camada'], d['campo']): d.get('apelido') for d in dicionario}
    for nome, frame in camadas.items():
        folha = livro.create_sheet(nome[:31])
        colunas = exportacao.colunas(frame)
        for j, coluna in enumerate(colunas, start=1):
            celula = folha.cell(row=1, column=j, value=coluna)
            celula.font = Font(bold=True)
            dica = apelidos.get((nome, coluna))
            if dica:
                from openpyxl.comments import Comment
                celula.comment = Comment(_texto_xlsx(dica), 'SICARD')
            folha.column_dimensions[get_column_letter(j)].width = min(max(len(coluna) * 0.9, 12), 45)
        for i, linha in enumerate(exportacao.registros(frame), start=2):
            for j, coluna in enumerate(colunas, start=1):
                folha.cell(row=i, column=j, value=_texto_xlsx(linha.get(coluna)))
        folha.freeze_panes = 'A2'
    folha = livro.create_sheet('dicionario_campos')
    folha.append(COLUNAS_DICIONARIO)
    for celula in folha[1]:
        celula.font = Font(bold=True)
    for item in dicionario:
        folha.append([_texto_xlsx(item.get(c)) for c in COLUNAS_DICIONARIO])
    livro.save(path)


def montar_pacote(camadas: dict, entrada, dicionario: list[dict], configuracao: dict,
                  nome_saida: str, finalidades=None, validacao=None, preservar_geometrias=False, incluir_entrada=True) -> tuple[bytes, str, list[dict]]:
    """Escreve os arquivos, confere cada um e devolve (zip, nome do zip, manifesto)."""
    arquivos = nomes(nome_saida, camadas, finalidades)
    if not incluir_entrada:
        arquivos['gpkg'] = (arquivos['gpkg'][0], 'GeoPackage de resultados e finalidades; entrada local temporária não incluída')
    with tempfile.TemporaryDirectory(prefix='sicard_enriquecimento_') as temporaria:
        pasta = Path(temporaria)
        escrever_gpkg(camadas, entrada, dicionario, pasta / arquivos['gpkg'][0], finalidades, preservar_geometrias, incluir_entrada)
        for nome, frame in camadas.items():
            exportacao.escrever_csv(frame, pasta / arquivos[f'csv_{nome}'][0])
        for chave, item in (finalidades or {}).items():
            for nome, frame in item['camadas'].items():
                exportacao.escrever_csv(frame, pasta / arquivos[f'csv_{chave}_{nome}'][0])
        (pasta / arquivos['validacao'][0]).write_text(
            json.dumps(validacao or {}, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
        escrever_xlsx(camadas, dicionario, pasta / arquivos['xlsx'][0])
        escrever_dicionario_csv(dicionario, pasta / arquivos['csv_dicionario'][0])
        (pasta / arquivos['configuracao'][0]).write_text(
            json.dumps(configuracao, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
        manifesto = []
        memoria = io.BytesIO()
        with zipfile.ZipFile(memoria, 'w', zipfile.ZIP_DEFLATED) as pacote:
            for chave, (nome, descricao) in arquivos.items():
                dados = (pasta / nome).read_bytes()
                if not dados:
                    raise ValueError(f'O arquivo {nome} do pacote saiu vazio.')
                pacote.writestr(nome, dados)
                manifesto.append({'chave': chave, 'nome': nome, 'descricao': descricao,
                                  'tamanho_bytes': len(dados), 'sha256': sha256(dados).hexdigest()})
    return memoria.getvalue(), f"{Path(arquivos['gpkg'][0]).stem}.zip", manifesto


def arquivos_analiticos(snapshot):
    from api.services.extracao_resultados_territoriais import consultar
    data = consultar(snapshot,tamanho=max(1,sum(len(e['feicoes']) for e in snapshot['entradas'])))
    stream=io.StringIO(); writer=csv.writer(stream,delimiter=';')
    writer.writerow(['camada_demanda','codigo_demanda','categoria','camada_base','elemento','pontos','comprimento_m','area_m2','perimetro_m','percentual_demanda','pontos_por_categoria','atributos_base'])
    for row in data['linhas']:
        for aid in row['areas']:
            area=data['areas'][aid]; m=row.get('relacoes',{}).get(aid,{})
            writer.writerow([exportacao.safe(v) for v in [row['entrada'],row.get('identificador'),area['categoria'],area['base'],area.get('nome',area['fid']),
                m.get('pontos'),m.get('comprimento_m'),m.get('area_m2'),m.get('perimetro_m'),m.get('percentual_entrada'),
                json.dumps(m.get('por_categoria',{}),ensure_ascii=False),json.dumps(area['atributos'],ensure_ascii=False,default=str)]])
    return {'analise_descritiva.json':json.dumps(data,ensure_ascii=False,default=str).encode(),
            'relacoes_por_demanda.csv':stream.getvalue().encode('utf-8-sig')}


def montar_lote(saida, configuracao, nome_saida):
    """ZIP por entrada e ZIP geral; tabela unificada sempre derivada das saídas."""
    import geopandas as gpd
    memoria = io.BytesIO()
    manifesto = []
    def adicionar(pacote, chave, nome, dados, descricao):
        pacote.writestr(nome, dados)
        manifesto.append({'chave':chave,'nome':nome,'descricao':descricao,'tamanho_bytes':len(dados),'sha256':sha256(dados).hexdigest()})
    with zipfile.ZipFile(memoria,'w',zipfile.ZIP_DEFLATED) as geral:
        for item in saida['individuais']:
            camadas = {n:saida['camadas'][n] for n in item['camadas']}
            dic = [d for d in saida['dicionario'] if d['camada'] in camadas]
            # Camada auxiliar contém somente elementos vinculados, nunca a base integral.
            from api.services.extracao_saida_analitica import snapshot_saida
            snapshot = snapshot_saida(camadas)
            areas = [a for a in snapshot['areas'].values() if a.get('geometria')]
            if areas:
                camadas = {**camadas, 'elementos_relacionados':gpd.GeoDataFrame.from_features([
                    {'type':'Feature','geometry':a['geometria'],'properties':{'elemento_id':a['id'],'base':a['base'],
                     'categoria':a['categoria'],'atributos':json.dumps(a['atributos'],ensure_ascii=False,default=str)}} for a in areas],crs=4674)}
            cfg = {**configuracao,'entrada':{k:v for k,v in item['entrada'].items() if k != 'frame'},'operacao':item['operacao']}
            fins = {k:{**f,'camadas':{n:df for n,df in f['camadas'].items() if n in item['camadas']}}
                    for k,f in saida['finalidades'].items() if any(n in item['camadas'] for n in f['camadas'])}
            conteudo, nome, _ = montar_pacote(camadas,item['entrada']['frame'],dic,cfg,item['nome_saida'],
                finalidades=fins,validacao=item['validacao'],preservar_geometrias=True,incluir_entrada=False)
            memoria_individual = io.BytesIO(conteudo)
            with zipfile.ZipFile(memoria_individual,'a',zipfile.ZIP_DEFLATED) as individual:
                for arquivo,dados in arquivos_analiticos(snapshot).items(): individual.writestr(arquivo,dados)
            conteudo = memoria_individual.getvalue()
            adicionar(geral,item['chave'],f"{item['chave']}/{nome}",conteudo,'Pacote individual: '+item['nome'])
            with zipfile.ZipFile(io.BytesIO(conteudo)) as individual:
                for arq in individual.namelist():
                    if arq.endswith('.gpkg'):
                        adicionar(geral,item['chave']+'_gpkg',f"geopackages/{item['chave']}_{arq}",individual.read(arq),'GeoPackage individual')
        with tempfile.TemporaryDirectory(prefix='sicard_lote_') as tmp:
            path = Path(tmp)
            escrever_xlsx(saida['camadas'],saida['dicionario'],path/'tabelas_unificadas.xlsx')
            adicionar(geral,'xlsx','unificado/tabelas.xlsx',(path/'tabelas_unificadas.xlsx').read_bytes(),'Tabelas de todas as saídas')
            snapshot = snapshot_saida(saida['camadas'])
            from api.services.extracao_resultados_territoriais import consultar
            for arquivo,dados in arquivos_analiticos(snapshot).items():
                adicionar(geral,'analitico_'+arquivo,'unificado/'+arquivo,dados,'Análise descritiva por demanda e elemento')
            ranking = consultar(snapshot,tamanho=max(1,sum(len(e['feicoes']) for e in snapshot['entradas'])))
            adicionar(geral,'analise','unificado/analise.json',json.dumps(ranking,ensure_ascii=False,default=str).encode(),'Análise descritiva unificada')
            columns = ['entrada','demanda','restricoes','riscos']
            stream=io.StringIO(); writer=csv.writer(stream,delimiter=';'); writer.writerow(columns)
            for row in ranking['linhas']:
                writer.writerow([exportacao.safe(row['entrada']),exportacao.safe(row.get('identificador')),row.get('contagens',{}).get('restricao',0),row.get('contagens',{}).get('risco',0)])
            adicionar(geral,'ranking','unificado/demandas.csv',stream.getvalue().encode('utf-8-sig'),'Comparação descritiva das demandas')
        adicionar(geral,'configuracao','unificado/configuracao.json',json.dumps(configuracao,ensure_ascii=False,default=str).encode(),'Procedência do lote')
    return memoria.getvalue(),f"{apelido(nome_saida,60) or 'extracao'}_lote.zip",manifesto
