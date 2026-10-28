"""Superfícies preparadas da Fase 2; valores persistidos independem dos pesos AHP."""
from __future__ import annotations

import io
import json
import math
from uuid import uuid4

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from psycopg.types.json import Jsonb

from api.db.connection import get_connection
from api.repositories import camada_geoespacial_repository as camadas
from api.services import ahp_engine, municipal_layer

DIMENSOES = (
    'Técnica', 'Financeiro', 'Econômica', 'Social', 'Segurança', 'Ambiental',
    'Territorial', 'Institucional', 'Risco', 'Fundiário-patrimonial',
    'Infraestrutura e servidões', 'Jurídico-institucional',
)


def gerar_conceitos(variavel: str, relacao: str) -> dict:
    label = variavel.strip().rstrip('.')
    if not label or relacao not in ('positiva', 'negativa'):
        raise ValueError('Informe uma variável descritiva e sua relação.')
    # A propriedade mantém o significado do catálogo, sem inferir benefício causal.
    return {
        'criterio': f'{label} (favorabilidade territorial)',
        'fenomeno_avaliado': 'favorabilidade territorial',
        'premissa': f"Quanto maior o valor de {label}, {'maior' if relacao == 'positiva' else 'menor'} a favorabilidade territorial.",
    }


def reescalonar(serie: pd.Series, relacao: str) -> tuple[pd.Series, dict]:
    if relacao not in ('positiva', 'negativa'):
        raise ValueError('Relação inválida.')
    values = pd.to_numeric(serie, errors='raise').astype(float)
    if np.isinf(values.to_numpy()).any():
        raise ValueError('A variável contém valores infinitos.')
    valid = values.dropna()
    if valid.empty:
        raise ValueError('A variável não possui observações numéricas para reescalonar.')
    low, high = float(valid.min()), float(valid.max())
    if low == high:
        raise ValueError('A variável é constante: a transformação linear mínimo–máximo é indefinida (denominador zero). Selecione uma variável com variação real.')
    scaled = (values - low) / (high - low)
    if relacao == 'negativa':
        scaled = 1 - scaled
    return scaled, {'metodo': 'linear_min_max', 'minimo': low, 'maximo': high,
                    'relacao': relacao, 'ausentes': int(values.isna().sum())}


def _detalhe(item):
    value = item.get('detail') or {}
    return json.loads(value) if isinstance(value, str) else value


def cobertura_comprovada(item, territorialidade, catalogo, visitados=None):
    """Cobertura do levantamento, nunca a contagem de valores numéricos."""
    detail = _detalhe(item)
    if detail.get('cobertura_territorial_integral') is True:
        return True
    seen = set(visitados or ())
    if item['id'] in seen:
        raise ValueError('Dependência circular na cobertura das variáveis combinadas.')
    seen.add(item['id'])
    refs = detail.get('atributos_origem') or []
    if refs:
        return all(ref in catalogo and cobertura_comprovada(
            catalogo[ref], territorialidade, catalogo, seen) for ref in refs)
    if territorialidade == 'setores_2022':
        # O importador registra o universo de linhas, inclusive sigilos e vazios.
        return detail.get('registros_fonte_sp') == 103319 and detail.get('setores_sem_linha') == 0
    # O Censo municipal tem âmbito de levantamento estadual; falta de observação
    # por categoria não é uma área excluída do levantamento.
    if item.get('source') == 'IBGE · Censo 2022':
        return True
    return detail.get('municipios_levantados') == 645


def catalogo(territorialidade):
    source = municipal_layer.backend(territorialidade)
    items = source.catalog()
    by_id = {i['id']: i for i in items}
    return {'attributes': [i for i in items if cobertura_comprovada(i, territorialidade, by_id)],
            'dimensoes': list(DIMENSOES), 'territorialidade': territorialidade,
            'featureCount': source.MUNICIPIOS}


def listar(hierarquizacao_id, user):
    with get_connection() as conn:
        return [dict(row) for row in conn.execute(
            '''SELECT id,nome,hierarquizacao_id,territorialidade,criado_em,calculo
               FROM geoprocessamento.superficie_criterios
               WHERE hierarquizacao_id=%s AND criado_por=%s ORDER BY criado_em DESC''',
            (hierarquizacao_id, user.id)).fetchall()]


def obter(ident, user):
    with get_connection() as conn:
        row = conn.execute('SELECT * FROM geoprocessamento.superficie_criterios WHERE id=%s AND criado_por=%s',
                           (ident, user.id)).fetchone()
    if not row:
        raise ValueError('Superfície não encontrada ou indisponível para este usuário.')
    return dict(row)


def carregar(row, indice=False):
    ident = row['camada_indice_id'] if indice else row['camada_id']
    if not ident:
        raise ValueError('O índice ainda não foi calculado.')
    loaded = camadas.carregar_vetor_bruto(ident)
    if loaded is None:
        raise ValueError('A camada salva não está disponível no acervo.')
    return loaded[0]


def preparar_frame(source, items, configuracoes):
    frame = source.layer(items)
    by_id = {i['id']: i for i in items}
    campos = municipal_layer.dados.nomes_exportados(items, 'gpkg')
    cadastrais = [c for c in frame.columns if c not in {i['field'] for i in items} and c != frame.geometry.name]
    output = frame[cadastrais + [frame.geometry.name]].copy()
    criterios = []
    for n, config in enumerate(configuracoes, 1):
        item = by_id[config['atributo_id']]
        if config['dimensao'] not in DIMENSOES:
            raise ValueError('Dimensão de critério inválida.')
        key = f'c{n:03d}'
        original = campos[item['field']]
        rescaled = original[:48] + '_reescalonado'
        if original in output or rescaled in output:
            raise ValueError('Colisão de nomes nas colunas da superfície.')
        values, transform = reescalonar(frame[item['field']], config['relacao'])
        conceitos = gerar_conceitos(item['label'], config['relacao'])
        criterios.append({'id': key, 'atributo_id': item['id'], 'variavel': item['label'],
                          **conceitos, 'dimensao': config['dimensao'], 'relacao': config['relacao'],
                          'campo_original': original, 'campo_reescalonado': rescaled,
                          'transformacao': transform, 'fonte': item['source'], 'ano': item['year'],
                          'unidade': item.get('unit'), 'url': item.get('url')})
        output[original] = frame[item['field']]
        output[rescaled] = values
    # Uma coluna para cada conceito, com mapeamento explícito pelo ID do critério.
    for name in ('premissa', 'criterio', 'dimensao'):
        output[name] = json.dumps({c['id']: c[name] for c in criterios}, ensure_ascii=False)
    pairs = [field for c in criterios for field in (c['campo_original'], c['campo_reescalonado'])]
    output = output[cadastrais + ['premissa', 'criterio', 'dimensao'] + pairs + [output.geometry.name]]
    return output, criterios


def salvar(payload, user):
    hierarchy = str(payload['hierarquizacao_id'])
    with get_connection() as conn:
        if not conn.execute('SELECT 1 FROM hierarquizacao_demandas.hierarquizacao_portfolio WHERE id=%s',
                            (hierarchy,)).fetchone():
            raise ValueError('Selecione uma hierarquização cadastrada.')
    configs = payload['criterios']
    ids = [c['atributo_id'] for c in configs]
    if len(ids) != len(set(ids)):
        raise ValueError('Uma variável não pode aparecer duas vezes na mesma superfície.')
    territory = payload['territorialidade']
    source = municipal_layer.backend(territory)
    allowed = {i['id']: i for i in catalogo(territory)['attributes']}
    if any(ident not in allowed for ident in ids):
        raise ValueError('Uma variável selecionada não possui cobertura territorial integral comprovada.')
    items = [allowed[ident] for ident in ids]
    frame, criteria = preparar_frame(source, items, configs)
    ident = str(uuid4())
    resource = 'superficie_' + ident
    camadas.salvar_vetor(recurso_id=resource, nome=payload['nome'], origem='OP-SUPERFICIE-CRITERIOS',
                         gdf=frame, metadados={'superficie_id': ident, 'hierarquizacao_id': hierarchy,
                                              'criterios': criteria, 'territorialidade': territory,
                                              'sem_indice': True}, preservar_geometrias=True)
    with get_connection() as conn:
        conn.execute('''INSERT INTO geoprocessamento.superficie_criterios
            (id,hierarquizacao_id,criado_por,nome,territorialidade,criterios,camada_id)
            VALUES (%s,%s,%s,%s,%s,%s,%s)''',
            (ident, hierarchy, user.id, payload['nome'], territory, Jsonb(criteria), resource))
        conn.commit()
    return obter(ident, user)


def pesos_matriz(matrix, n):
    if len(matrix) != n or any(len(row) != n for row in matrix):
        raise ValueError('A matriz deve corresponder exatamente aos critérios da superfície.')
    for i in range(n):
        for j in range(n):
            value = matrix[i][j]
            if not math.isfinite(value) or not 1/9 <= value <= 9:
                raise ValueError('Informe comparações finitas entre 1/9 e 9.')
            if i == j and not math.isclose(value, 1):
                raise ValueError('A diagonal da matriz deve ser igual a 1.')
            if not math.isclose(value * matrix[j][i], 1, rel_tol=1e-8):
                raise ValueError('A matriz de comparação deve ser recíproca.')
    result = ahp_engine.analyze_matrix(matrix)
    if result['CR'] > 0.1 + 1e-10:
        raise ValueError(f"Razão de consistência {result['CR']:.3f} acima de 0,10. Revise as comparações.")
    return result


def compor(frame, criterios, pesos):
    if len(pesos) != len(criterios) or any(not math.isfinite(w) or w <= 0 for w in pesos):
        raise ValueError('Informe um peso positivo e finito para cada critério.')
    total = sum(pesos)
    if not math.isclose(total, 1, rel_tol=1e-6, abs_tol=1e-8):
        raise ValueError('Os pesos devem somar 1.')
    values = frame[[c['campo_reescalonado'] for c in criterios]].astype(float)
    finite = values.to_numpy().ravel()
    finite = finite[~np.isnan(finite)]
    if finite.size and (not np.isfinite(finite).all() or finite.min() < 0 or finite.max() > 1):
        raise ValueError('A camada salva contém valores reescalonados fora de 0 a 1.')
    # Zero observado é válido. NoData/sigilo não se transforma em zero.
    return values.mul(pesos).sum(axis=1, min_count=len(criterios))


def calcular(ident, payload, user):
    if payload.get('matriz') is not None or payload.get('ambiente_id'):
        raise ValueError('Julgamentos AHP são exclusivos da Fase 3. Na Fase 2, utilize média simples ou pesos explicitamente informados.')
    row = obter(ident, user)
    criteria = row['criterios']
    method = payload['operador']
    audit = {'operador': method, 'criterios_ids': [c['id'] for c in criteria],
             'ausentes': 'indice_nulo_sem_redistribuicao'}
    if method == 'media_simples':
        weights = [1 / len(criteria)] * len(criteria)
    elif payload.get('ambiente_id'):
        from api.services.superficie_ahp import pesos_consolidados
        weights = pesos_consolidados(row, payload['ambiente_id'])
        audit['ambiente_id'] = payload['ambiente_id']
    elif payload.get('matriz') is not None:
        result = pesos_matriz(payload['matriz'], len(criteria))
        weights = result['weights']
        audit.update(matriz=payload['matriz'], consistencia=result)
    else:
        supplied = payload.get('pesos') or {}
        if set(supplied) != {c['id'] for c in criteria}:
            raise ValueError('Os pesos devem identificar exatamente os critérios desta superfície.')
        weights = [supplied[c['id']] for c in criteria]
        audit['origem'] = 'pesos_informados'
    frame = carregar(row)
    frame['indice_favorabilidade'] = compor(frame, criteria, weights)
    audit['pesos'] = dict(zip([c['id'] for c in criteria], weights))
    audit['sem_indice'] = int(frame['indice_favorabilidade'].isna().sum())
    resource = 'indice_' + str(uuid4())
    camadas.salvar_vetor(recurso_id=resource, nome=row['nome'] + ' — índice',
                         origem='OP-SUPERFICIE-FAVORABILIDADE', gdf=frame,
                         metadados={'superficie_id': str(row['id']), 'criterios': criteria,
                                    'hierarquizacao_id': str(row['hierarquizacao_id']), 'calculo': audit},
                         preservar_geometrias=True)
    with get_connection() as conn:
        conn.execute('''UPDATE geoprocessamento.superficie_criterios
            SET calculo=%s,camada_indice_id=%s,atualizado_em=now() WHERE id=%s AND criado_por=%s''',
            (Jsonb(audit), resource, ident, user.id))
        conn.commit()
    return obter(ident, user)


def visualizar(ident, user, campo):
    row = obter(ident, user)
    fields = {f for c in row['criterios'] for f in (c['campo_original'], c['campo_reescalonado'])}
    if campo == 'indice_favorabilidade' and row['camada_indice_id']:
        frame = carregar(row, indice=True)
    else:
        if campo not in fields:
            raise ValueError('Campo não pertence à camada de critérios salva.')
        frame = carregar(row)
    # Simplificação exclusiva de exibição; o produto persistido permanece intacto.
    display = frame[[campo, frame.geometry.name]].copy()
    display.geometry = display.geometry.simplify(0.0002, preserve_topology=True)
    return json.loads(display.to_json(drop_id=True, na='null'))


def matriz_excel(row):
    from api.path_policy import project_path
    book = load_workbook(project_path('documentacao/matrizes/Modelo_Matriz_Criterios_Premissas_SLT.xlsx'))
    sheet = book['Critérios']
    sheet.delete_rows(2, max(sheet.max_row - 1, 1))
    matrix_sheet = book['Matriz Crit Premissas v3']
    matrix_sheet.delete_rows(2, max(matrix_sheet.max_row - 1, 1))
    for n, c in enumerate(row['criterios'], 1):
        sheet.append([n, c['criterio'], c['dimensao']])
        matrix_sheet.append([
            c['dimensao'], c['criterio'], c['premissa'], 'Favorabilidade territorial — município e rede rodoviária',
            None, 'Fase 2', c.get('fonte'), c['variavel'], c.get('unidade'),
            'Transformação linear mínimo–máximo',
            '↑ Positiva' if c['relacao'] == 'positiva' else '↓ Negativa',
            c['campo_reescalonado'], c.get('url'), 'Não',
        ])
    sheet = book.create_sheet('Comparação pareada')
    criteria = row['criterios']
    sheet.append(['Critério'] + [c['criterio'] for c in criteria])
    for i, c in enumerate(criteria):
        sheet.append([c['criterio']] + [1 if i == j else None for j in range(len(criteria))])
    sheet = book.create_sheet('Vínculo')
    sheet.append(['Superfície', str(row['id'])])
    sheet.append(['Hierarquização', str(row['hierarquizacao_id'])])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def exportar(ident, user, indice=False):
    from api.services import saidas_storage
    row = obter(ident, user)
    resource = row['camada_indice_id'] if indice else row['camada_id']
    if not resource:
        raise ValueError('O índice ainda não foi calculado.')
    loaded = camadas.carregar_vetor_bruto(resource)
    if loaded is None:
        raise ValueError('Camada salva não disponível no acervo.')
    info = loaded[1]
    return saidas_storage.conferir(info['storage_caminho'], (info.get('metadados') or {}).get('sha256'))
