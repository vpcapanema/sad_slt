"""Tabelas e guia da Fase 2 a partir dos resultados persistidos, sem novo cálculo espacial."""
from __future__ import annotations

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment

from api.services import elegibilidade_sintese as tabelas

PAPEIS = {'grade': 'Grade', 'rede': 'Rede'}
NOTA = (
    'Cada linha de detalhe identifica uma combinação execução × demanda × superfície × feição. '
    'Os índices e classes são os da fonte, sem média, soma ou ponderação por área. '
    'Zero válido não é ausência. Índices ausentes, nulos ou inválidos permanecem sem valor numérico '
    'e têm estado explícito. Uma linha sem interseção não recebe índice nem classe. '
    'As áreas são informação técnica da interseção, não índices; sobreposições internas podem '
    'produzir percentuais acima de 100%. As tabelas não reproduzem as estatísticas exploratórias '
    'selecionadas na página e não constituem ranking ou homologação.'
    ' Para raster, o detalhe preserva os valores das células; a tabela Síntese Raster informa '
    'a média ponderada pela área efetiva de interseção e a cobertura válida de cada demanda. '
    'Zero entra na média; NoData e áreas externas ficam fora do denominador. Sem cobertura válida, '
    'o índice é nulo. Não se atribuem classes aos rasters contínuos.'
)
COLUNAS_RASTER = ['Execução', 'Conjunto de demandas', 'Chave da demanda', 'Demanda',
                  'ID da superfície', 'Papel', 'Métrica', 'Índice ponderado',
                  'Área válida (m²)', 'Cobertura válida (%)', 'Área sem dados (m²)', 'Células válidas']


def sintese_raster(modelo: dict) -> list[list]:
    return [[entrada['id'], entrada['nome'], demanda['chave'], demanda['nome'],
             r['base_id'], PAPEIS[r['papel']], r['metrica'], r['indice'], r['area_valida_m2'],
             r['cobertura_percentual'], r['area_sem_dados_m2'], r['celulas_validas']]
            for entrada in modelo['entradas'] for demanda in entrada['linhas'] for r in demanda.get('rasters', [])]
COLUNAS_DETALHE = [
    'Execução', 'Conjunto de demandas', 'Chave da demanda', 'Identificador da demanda', 'Demanda',
    'ID da superfície', 'Papel', 'Superfície', 'ID da feição', 'Posição da feição', 'Feição',
    'Situação espacial', 'Campo do índice dimensionado', 'Índice dimensionado', 'Estado do índice dimensionado',
    'Campo da classe', 'Classe original', 'Estado da classe',
    'Campo do índice sem _dim', 'Índice sem _dim', 'Estado do índice sem _dim',
    'Área da demanda (m²)', 'Área da interseção (m²)', '% da área da demanda',
]
COLUNAS_RESUMO = [
    'Execução', 'Conjunto de demandas', 'ID da superfície', 'Papel', 'Superfície',
    'Demandas analisadas', 'Demandas com interseção', 'Demandas sem interseção',
    'Combinações demanda × feição', 'Índices válidos', 'Índices ausentes', 'Índices nulos',
    'Índices inválidos', 'Contatos sem área',
]


def detalhe(modelo: dict) -> list[list]:
    linhas = []
    for entrada in modelo['entradas']:
        for demanda in entrada['linhas']:
            for superficie in entrada['campos']:
                itens = [item for item in demanda['sobreposicoes'] if item['base_id'] == superficie['id']]
                comuns = [entrada['id'], entrada['nome'], demanda['chave'], demanda['id'], demanda['nome'],
                          superficie['id'], PAPEIS[superficie['papel']], superficie['nome']]
                if not itens:
                    linhas.append([*comuns, None, None, None, 'Sem interseção',
                                   None, None, 'sem_intersecao', None, None, 'sem_intersecao',
                                   None, None, 'sem_intersecao', demanda['area_total_m2'], None, None])
                for item in itens:
                    valor = item['favorabilidade']
                    linhas.append([*comuns, item['id'], item['identificador'], item['nome'],
                                   'Contato sem área' if item['area_m2'] == 0 else 'Interseção com área',
                                   valor['campo_indice'], valor['indice'], valor['estado_indice'],
                                   valor['campo_classe'], valor['classe'], valor['estado_classe'],
                                   valor['campo_original'], valor['indice_original'], valor['estado_original'],
                                   demanda['area_total_m2'], item['area_m2'], item['percentual']])
    return linhas


def resumo(modelo: dict) -> list[list]:
    linhas = []
    for entrada in modelo['entradas']:
        for superficie in entrada['campos']:
            itens = [item for demanda in entrada['linhas'] for item in demanda['sobreposicoes']
                     if item['base_id'] == superficie['id']]
            com = sum(any(item['base_id'] == superficie['id'] for item in demanda['sobreposicoes'])
                      for demanda in entrada['linhas'])
            estados = [item['favorabilidade']['estado_indice'] for item in itens]
            linhas.append([entrada['id'], entrada['nome'], superficie['id'], PAPEIS[superficie['papel']],
                           superficie['nome'], len(entrada['linhas']), com, len(entrada['linhas']) - com,
                           len(itens), *[estados.count(estado) for estado in ('valido', 'ausente', 'nulo', 'invalido')],
                           sum(item['area_m2'] == 0 for item in itens)])
    return linhas


def _csv(colunas: list[str], linhas: list[list]) -> bytes:
    memoria = StringIO()
    writer = csv.writer(memoria, delimiter=';', lineterminator='\r\n')
    writer.writerow(colunas)
    for linha in linhas:
        # repr conserva a precisao do float, sem o arredondamento de exibicao do PDF.
        writer.writerow([repr(valor).replace('.', ',') if isinstance(valor, float)
                         else tabelas._celula_csv(valor) for valor in linha])
    return memoria.getvalue().encode('utf-8-sig')


def xlsx(modelo: dict, detalhes: list[list], resumos: list[list]) -> bytes:
    livro = Workbook()
    tabelas._folha(livro.active, 'Índices e Classes', COLUNAS_DETALHE, detalhes,
                  [25, 30, 18, 24, 38, 28, 16, 36, 28, 18, 38, 25, 30, 20, 25, 30, 26, 22,
                   30, 20, 25, 24, 24, 24])
    for linha in livro.active.iter_rows(min_row=2):
        for indice in (13, 19, 21, 22, 23):
            linha[indice].number_format = '0.###############'
    tabelas._folha(livro.create_sheet(), 'Cobertura por Superfície', COLUNAS_RESUMO, resumos,
                  [25, 30, 28, 16, 36] + [24] * 9)
    if any(c.get('tipo') == 'raster' for e in modelo['entradas'] for c in e['campos']):
        tabelas._folha(livro.create_sheet(), 'Síntese Raster', COLUNAS_RASTER, sintese_raster(modelo), [28] * len(COLUNAS_RASTER))
    tabelas._folha(livro.create_sheet(), 'Fontes e Nomes',
                  ['ID', 'Tipo', 'Papel', 'Nome revisado'],
                  [[fonte['id'], fonte['tipo'], fonte['papel'], fonte['nome']] for fonte in modelo['fontes']],
                  [36, 20, 20, 50])
    folha = livro.create_sheet('Como Ler')
    for texto in (NOTA, modelo['metodologia'],
                  'Estados: valido = valor utilizável; ausente = campo não existe na fonte; '
                  'nulo = campo sem valor; invalido = tipo ou domínio incompatível; '
                  'sem_intersecao = nenhuma feição encontrada nessa superfície para a demanda.',
                  'Chave da demanda é a posição na entrada, não um identificador global. Use-a junto da execução. '
                  'IDs de superfície e feição distinguem homônimos. Áreas estão em m²; percentuais em 0–100 '
                  '(não em fração). Índice dimensionado válido está em [0,1]; o índice sem _dim conserva '
                  'seu domínio de origem e não o substitui.',
                  'As tabelas servem para conferência e análise externa, não para reimportação ou atualização '
                  'automática. Para corrigir dados, ajuste a fonte e gere nova execução; preserve o pacote original.'):
        folha.append([texto])
    folha.column_dimensions['A'].width = 120
    for linha in folha:
        linha[0].alignment = Alignment(wrap_text=True, vertical='top')
        folha.row_dimensions[linha[0].row].height = 90
    memoria = BytesIO()
    livro.save(memoria)
    return memoria.getvalue()


def documentos(modelo: dict, relatorio: str, arquivos: list[str]) -> list[tuple[str, bytes]]:
    detalhes, resumos = detalhe(modelo), resumo(modelo)
    guia = [
        'RELATÓRIO E CAMADAS — FAVORABILIDADE TERRITORIAL', '',
        f'Entradas concluídas: {modelo["concluidas"]} de {modelo["previstas"]}.',
        'PACOTE PARCIAL: entradas não concluídas não recebem resultado.'
        if modelo['concluidas'] != modelo['previstas'] else 'Todas as entradas previstas foram concluídas.',
        '', 'ORDEM DE LEITURA',
        f'1. {relatorio}: alcance, cobertura por superfície, mapa e valores originais por demanda e feição.',
        '2. Tabelas/Favorabilidade Territorial.xlsx: detalhe editável, cobertura por superfície, fontes e guia.',
        '3. Tabelas/Índices e Classes.csv: o mesmo detalhe em formato aberto.',
        '4. Tabelas/Cobertura por Superfície.csv: contagens para conferir presença e qualidade dos índices, sem escore agregado.',
        '5. Camadas/NN/*.gpkg: produtos integrais para SIG, com a entrada preparada e a camada intersecoes, '
        'todos os atributos de origem, vínculos e metadados SICARD_FAVORABILIDADE.',
        '6. Metadados da Análise.json: execuções, nomes revisados e SHA-256 das camadas, PDF, tabelas e deste guia.',
        '', 'COMO LER', NOTA,
        'CSV: UTF-8 com BOM; separador ponto e vírgula; decimal vírgula; célula vazia não significa zero. '
        'Campos de estado distinguem ausente, nulo, invalido e sem_intersecao; valido indica valor utilizável. '
        'Percentuais são apresentados em 0–100, não em fração. XLSX preserva tipos numéricos; '
        'CSV preserva a precisão disponível sem o arredondamento de apresentação do PDF.',
        '', 'CONFERÊNCIA E EVENTUAIS AJUSTES',
        'Cruze execução + chave da demanda + ID da superfície + ID da feição, não apenas os nomes. '
        'Homônimos permanecem distintos. Os arquivos não são um formulário de reimportação. '
        'Alterar Excel ou CSV não atualiza o SICARD. Ajuste dados na fonte, reprocesse e compare '
        'com o pacote original preservado; não substitua ausências por zero nem reclassifique automaticamente.',
        'Confira cada assinatura SHA-256 antes de comparar arquivos. Não altere o manifesto para encobrir '
        'edições externas. A emissão do pacote não altera ranking, AHP ou homologação.',
        '', 'ARQUIVOS DE CAMADAS', *[f'- {nome}' for nome in arquivos],
    ]
    arquivos_raster = ([('Tabelas/Síntese Raster.csv', _csv(COLUNAS_RASTER, sintese_raster(modelo)))]
                       if any(c.get('tipo') == 'raster' for e in modelo['entradas'] for c in e['campos']) else [])
    return [
        ('Tabelas/Favorabilidade Territorial.xlsx', xlsx(modelo, detalhes, resumos)),
        ('Tabelas/Índices e Classes.csv', _csv(COLUNAS_DETALHE, detalhes)),
        ('Tabelas/Cobertura por Superfície.csv', _csv(COLUNAS_RESUMO, resumos)),
        ('LEIA-ME.txt', ('\r\n'.join(guia) + '\r\n').encode('utf-8-sig')),
        *arquivos_raster,
    ]
