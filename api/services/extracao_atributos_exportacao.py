"""Arquivos da mesma análise usada na tela; sem recalcular estatísticas.

Quem monta o pacote de cada extração é extracao_atributos_pacote: estas funções
só escrevem cada arquivo no caminho recebido.
"""
import csv
import json
from xml.sax.saxutils import escape

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

# Uma célula XLSX guarda no máximo 32.767 caracteres. O pacote é obrigatório, então
# o texto é cortado com aviso; o CSV do mesmo pacote leva o conteúdo completo.
LIMITE_CELULA_XLSX = 32767
AVISO_CORTE = ' … [texto completo no CSV do pacote]'
OPERACOES = {'intersection':'Interseção','identity':'Identidade'}


def data_hora(valor):
    """ISO em UTC -> data e hora de São Paulo; texto que não for data passa como veio."""
    from datetime import datetime
    from zoneinfo import ZoneInfo
    try:
        return datetime.fromisoformat(str(valor)).astimezone(ZoneInfo('America/Sao_Paulo')).strftime('%d/%m/%Y %H:%M:%S')
    except ValueError:
        return str(valor)


def medidas(value, dim):
    if value is None: return '-'
    def n(v,places): return f'{v:,.{places}f}'.translate(str.maketrans('.,',',.'))
    if dim == 1: return f'{n(value,3)} m / {n(value/1000,6)} km'
    if dim == 2: return f'{n(value/10000,6)} ha / {n(value/1000000,8)} km²'
    return '-'


def linhas(result):
    dim = result['dimensao_input']
    for category in result['categorias']:
        for layer in category['camadas']:
            for row in layer['ocorrencias']:
                if dim and (row.get('dimensao') != dim or row.get('medida_si',0)<=0): continue
                yield category,layer,row


def cells(result):
    dim = result['dimensao_input']
    units = ['Dentro'] if dim==0 else ['Metros','Quilômetros','Percentual'] if dim==1 else ['Hectares','Quilômetros quadrados','Percentual']
    yield ['Categoria','Camada','Feição input','Feição base',*units,'Atributos input','Atributos extraídos']
    for category,layer,row in linhas(result):
        value = row.get('medida_si',0)
        values = [row['dentro']] if dim==0 else [value,value/1000,row['percentual']] if dim==1 else [value/10000,value/1000000,row['percentual']]
        yield [category['nome'],layer['nome'],row['input_id'],row['feicao_base_id'],*values,
               json.dumps(row.get('atributos_input',{}),ensure_ascii=False),json.dumps(row['atributos'],ensure_ascii=False)]


def safe(value):
    # Evita que atributos externos sejam executados como fórmulas no Excel/CSV.
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r')):
        return "'"+value
    return value


def pdf(result,path):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Cell',fontName='Helvetica',fontSize=8,leading=11,wordWrap='CJK'))
    styles['BodyText'].fontSize=9; styles['BodyText'].leading=12
    styles['BodyText'].spaceBefore=2
    for name in ['Heading1','Heading2','Heading3','Heading4']:
        styles[name].spaceBefore=6;styles[name].spaceAfter=3
    styles['Title'].textColor=colors.HexColor('#173e59')
    def p(value,style='BodyText'): return Paragraph(escape(str(value)),styles[style])
    def table(rows,widths=None):
        item = Table([[p(v,'Cell') for v in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
        item.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e6eff5')),
                                 ('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.7,colors.HexColor('#42728b')),
                                 ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f5f7f8')]),
                                 ('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6)]))
        return item
    dim = result['dimensao_input']; summary = result['resumo']
    story = [p('SICARD | Extração de atributos','Title'),p(result['input_nome'],'Heading2'),
             p(f"Execução: {result['id']} | {data_hora(result['criado_em'])}"),
             p(f"Motor: {result['motor']} | GDAL {result['gdal']} | Operação: {OPERACOES.get(result['operacao'],result['operacao'])}"),
             p('Síntese da análise','Heading2'),
             table([['Categorias','Camadas intersectadas','Ocorrências','Parcela da entrada'],
                    [len(result['categorias']),summary['camadas_intersectadas'],summary['ocorrencias'],f"{summary['percentual']:.4f}%"]],[160]*4),
             Spacer(1,10),p('Total da entrada: '+(str(result['medida_total_input_si'])+' pontos' if dim==0 else medidas(result['medida_total_input_si'],dim))),
             p(result['metodologia_estatistica']),p(result['convencao_ids'])]
    for category in result['categorias']:
        story.extend([p(category['nome'],'Heading1'),p(category['conceito']),
                      p(f"Percentual da categoria: {category['resumo']['percentual']:.4f}%")])
        entities = [('Categoria consolidada',category),*((layer['nome'],layer) for layer in category['camadas'])]
        if dim:
            rows = [['Agrupamento','N','Média','Mediana','Mínimo','Máximo','Desvio padrão','Q1 / Q3']]
            for name,entity in entities:
                s=entity['estatisticas']
                rows.append([name,s['n'],*[medidas(s[k],dim) for k in ['media','mediana','minimo','maximo','desvio_padrao']],medidas(s['q1'],dim)+' / '+medidas(s['q3'],dim)])
            story.append(table(rows,[110,30,90,90,90,90,90,150]))
        else:
            story.append(table([['Agrupamento','Pontos dentro','Pontos fora'],*[[name,e['estatisticas']['pontos_dentro'],e['estatisticas']['pontos_fora']] for name,e in entities]],[350,180,210]))
        for layer in category['camadas']:
            story.append(p(layer['nome'],'Heading2'))
            r=layer['resumo']
            story.append(p(f"Ocorrências: {r['ocorrencias']} | {r['percentual']:.4f}% da entrada"))
            if dim: story.append(p('Extensão única: '+medidas(r['medida_unica_si'],dim)))
            for _,_,row in linhas({**result,'categorias':[{**category,'camadas':[layer]}]}):
                position = ('Dentro' if row['dentro'] else 'Fora') if dim==0 else medidas(row['medida_si'],dim)+f" | {row['percentual']:.4f}%"
                story.append(p(f"Input {row['input_id']} / base {row['feicao_base_id'] if row['feicao_base_id'] is not None else '-'}: {position}",'Heading3'))
                for label,values in [('Atributos da entrada',row.get('atributos_input',{})),('Atributos extraídos',row['atributos'])]:
                    story.append(p(label,'Heading4'))
                    if not values: story.append(p('Sem atributos nesta ocorrência.'))
                    for key,value in values.items():
                        text = f'{key}: '+(json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else str(value))
                        # Fluxo paginável para campos longos, sem truncar atributos.
                        story.extend(p(text[i:i+1000]) for i in range(0,len(text),1000))
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont('Helvetica',8)
        canvas.drawString(36,20,'SICARD - Extração espacial de atributos')
        canvas.drawRightString(805,20,f'Página {doc.page}');canvas.restoreState()
    SimpleDocTemplate(str(path),pagesize=landscape(A4),rightMargin=36,leftMargin=36,
                      topMargin=32,bottomMargin=34,title='Extração de atributos - SICARD').build(story,onFirstPage=footer,onLaterPages=footer)


def escrever_csv(result,path):
    with open(path,'w',newline='',encoding='utf-8-sig') as stream:
        csv.writer(stream,delimiter=';').writerows([[safe(v) for v in row] for row in cells(result)])


def celula_xlsx(value):
    value = safe(value)
    if isinstance(value,str) and len(value) > LIMITE_CELULA_XLSX:
        return value[:LIMITE_CELULA_XLSX-len(AVISO_CORTE)]+AVISO_CORTE
    return value


def escrever_xlsx(result,path):
    book=Workbook();sheet=book.active;sheet.title='Ocorrências'
    for row in cells(result):
        sheet.append([celula_xlsx(v) for v in row])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    summary=book.create_sheet('Síntese');summary.append(['Categoria','Camada','Ocorrências','Medida única SI','Percentual'])
    stats=book.create_sheet('Estatísticas');stats.append(['Categoria','Camada','Indicador','Valor'])
    for c in result['categorias']:
        for name,e in [('Categoria consolidada',c),*((l['nome'],l) for l in c['camadas'])]:
            r=e['resumo'];summary.append([celula_xlsx(c['nome']),celula_xlsx(name),r['ocorrencias'],r.get('medida_unica_si'),r['percentual']])
            for key,value in e['estatisticas'].items():stats.append([celula_xlsx(c['nome']),celula_xlsx(name),key,value])
    meta=book.create_sheet('Metodologia');meta.append([result['metodologia_estatistica']]);meta.append([result['convencao_ids']])
    book.save(path)
