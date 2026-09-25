"""Casos de borda isolados; não substituem a malha e o catálogo reais."""
import csv
import importlib.util
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from shapely.geometry import Polygon

# Carregamento explícito evita conflito com outros módulos chamados app.
spec = importlib.util.spec_from_file_location('municipal_standalone', Path(__file__).with_name('app.py'))
app = importlib.util.module_from_spec(spec)
spec.loader.exec_module(app)


class PackageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.data = Path(self.temp.name)
        frame = gpd.GeoDataFrame({'CD_MUN':['3500001','3500002','3500003'],
            'NM_MUN':['Água','=Texto literal','Sem dado'], 'SIGLA_UF':['SP']*3, 'AREA_KM2':[1.,2.,3.]},
            geometry=[Polygon([(x,0.12345678901234567),(x+1,0),(x+1,1),(x,0.12345678901234567)]) for x in range(3)],crs=4674)
        frame.to_file(self.data/'municipios.gpkg',layer='municipios',driver='GPKG')
        with sqlite3.connect(self.data/'catalog.sqlite') as conn:
            conn.executescript('CREATE TABLE attributes(id,source,year,theme,field,label,unit,url,detail,coverage); CREATE TABLE observations(attribute_id,municipality,value);')
            for ident in ['parcial','vazio']:
                conn.execute('INSERT INTO attributes VALUES (?,?,?,?,?,?,?,?,?,?)',
                             (ident,'Fonte teste',2022,'tema',ident,'Alias '+ident,'unidade','https://example.test','{}',1 if ident=='parcial' else 0))
            conn.executemany('INSERT INTO observations VALUES (?,?,?)',[
                ('parcial','3500001',0),('parcial','3500002',None),('parcial','3599999',99)])
        self.addCleanup(patch.stopall)
        patch.object(app,'DATA',self.data).start()

    def test_all_formats_and_companions(self):
        for fmt in ['fgb','gpkg','shp','geojson']:
            with self.subTest(format=fmt), tempfile.TemporaryDirectory() as temp:
                with zipfile.ZipFile(io.BytesIO(app.export_layer({'attributes':['parcial','vazio'],'format':fmt}))) as z:
                    self.assertIsNone(z.testzip())
                    z.extractall(temp)
                    names=set(z.namelist())
                root=Path(temp)
                for suffix in ['csv','txt','xlsx']:
                    self.assertIn('municipios_sp_atributos.'+suffix,names)
                self.assertIn('municipios_sp_relatorio_join.txt',names)
                manifest=json.loads((root/'metadados.json').read_text())
                self.assertEqual(set(manifest['arquivos']),names)
                self.assertEqual(manifest['join']['feicoes_entrada'],3)
                self.assertEqual(manifest['join']['campos']['parcial']['sem_observacao'],1)
                self.assertEqual(manifest['join']['campos']['parcial']['codigos_fora_da_malha'],['3599999'])
                actual=gpd.read_file(root/f'municipios_sp.{fmt}').set_index('CD_MUN').sort_index()
                fields={i['field']:i['export_field'] for i in manifest['attributes']}
                self.assertEqual(len(actual),3)
                self.assertEqual(actual.loc['3500001',fields['parcial']],0)
                self.assertTrue(actual[fields['vazio']].isna().all())
                self.assertTrue(actual.loc[['3500002','3500003'],fields['parcial']].isna().all())
                for ext,sep in [('csv',','),('txt','\t')]:
                    tab=pd.read_csv(root/f'municipios_sp_atributos.{ext}',sep=sep,dtype={'CD_MUN':str}).set_index('CD_MUN').sort_index()
                    pd.testing.assert_frame_equal(tab, pd.DataFrame(actual.drop(columns='geometry')),check_dtype=False)
                workbook=load_workbook(root/'municipios_sp_atributos.xlsx',read_only=True)
                rows=list(workbook.active.iter_rows(max_col=len(actual.columns),values_only=True))
                self.assertEqual(len(rows),4)
                data={r[0]:dict(zip(rows[0],r)) for r in rows[1:]}
                self.assertIsNone(data['3500003'][fields['parcial']])
                self.assertEqual(data['3500002']['NM_MUN'],'=Texto literal')
                workbook.close()
                with (root/'dicionario.csv').open(encoding='utf-8-sig') as f:
                    glossary=list(csv.DictReader(f))
                self.assertEqual(len(glossary),6)
                self.assertTrue(all(r['campo_bruto'] and r['alias'] for r in glossary))
                if fmt=='shp':
                    self.assertTrue({'municipios_sp'+e for e in ['.shp','.shx','.dbf','.prj','.cpg']}<=names)

    def test_duplicate_observations_refused(self):
        with sqlite3.connect(self.data/'catalog.sqlite') as conn:
            conn.execute("INSERT INTO observations VALUES ('parcial','3500001',5)")
        with self.assertRaisesRegex(ValueError,'repetidas'):
            app.layer(app.selection(['parcial']))

    def test_reopened_layer_must_preserve_codes(self):
        frame=app.layer(app.selection(['parcial']))
        path=self.data/'truncada.geojson'
        frame.iloc[:2].to_file(path,driver='GeoJSON')
        with self.assertRaisesRegex(ValueError,'quantidade'):
            app.export_support.reopen_and_validate(path,frame)


if __name__=='__main__':
    unittest.main()
