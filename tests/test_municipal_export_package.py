"""Contrato do exportador SICARD com uma malha isolada de três municípios."""

from contextlib import contextmanager
import csv
import io
import json
import zipfile
from types import SimpleNamespace
from uuid import uuid4

import geopandas as gpd
import pandas as pd
import pytest
from openpyxl import load_workbook
from shapely.geometry import MultiPolygon, Polygon

from api.services import base_municipal as dados
from api.routers.municipal_layer import Selecao


@pytest.fixture
def storage_api(monkeypatch):
    from api import storage_policy
    from api.services import storage_remoto
    import tempfile

    def forbid(*args, **kwargs):
        raise AssertionError('A geração tentou criar temporários no disco local.')

    for key in ('SICARD_STORAGE_WRITE_DIR', 'SICARD_STORAGE_TEMP_DIR'):
        monkeypatch.delenv(key, raising=False)
    monkeypatch.setattr(storage_policy, 'raiz_temporaria', forbid)
    monkeypatch.setattr(tempfile, 'mkstemp', forbid)
    monkeypatch.setattr(tempfile, 'NamedTemporaryFile', forbid)
    files = {}
    monkeypatch.setattr(storage_remoto, 'listar', lambda *a, **kw: [])
    monkeypatch.setattr(storage_remoto, 'criar_pasta', lambda path: None)
    monkeypatch.setattr(storage_remoto, 'apagar_pasta', lambda path: None)
    monkeypatch.setattr(storage_remoto, 'enviar_bytes', lambda path, data: files.update({path: data}))
    monkeypatch.setattr(storage_remoto, 'baixar', lambda path: files[path])
    monkeypatch.setattr(storage_remoto, 'apagar_arquivo', lambda path: files.pop(path, None))
    yield files
    assert not any(path.startswith('temporarios-sicard/') for path in files)


@pytest.fixture(params=['montado', 'api'])
def inputs(monkeypatch, storage_sintetico, request):
    if request.param == 'api':
        request.getfixturevalue('storage_api')
    frame = gpd.GeoDataFrame(
        {
            'CD_MUN': ['3500001', '3500002', '3500003'],
            'NM_MUN': ['Um', 'Dois', 'Três'],
            'SIGLA_UF': ['SP'] * 3,
            'AREA_KM2': [1., 2., 3.],
        },
        geometry=[
            MultiPolygon([
                Polygon([
                    (x, 0.12345678901234567),
                    (x + 1, 0),
                    (x + 1, 1),
                    (x, 0.12345678901234567),
                ])
            ])
            for x in range(3)
        ],
        crs=4674,
    ).set_index('CD_MUN')
    items = [
        dict(
            id=f,
            field=f,
            label='Indicador ' + f,
            source='Fonte de teste',
            theme='economia',
            year=2022,
            unit='unidade',
            url='https://example.test',
            detail='{}',
            coverage=1 if f == 'parcial' else 0,
        )
        for f in ['parcial', 'vazio']
    ]
    rows = [
        dict(
            atributo_id='parcial',
            municipios=['3500001', '3500002', '3599999'],
            valores=[0, None, 10],
        )
    ]

    class Connection:
        def execute(self, *args):
            return self

        def fetchall(self):
            return rows

    @contextmanager
    def connection():
        yield Connection()

    monkeypatch.setattr(dados, 'malha', frame.copy)
    monkeypatch.setattr(dados, 'catalog', lambda: items)
    monkeypatch.setattr(dados, 'get_connection', connection)
    monkeypatch.setenv('SLT_MUNICIPAL_ORCAMENTO_MB', '99999')
    return items, frame


@pytest.mark.parametrize('fmt', ['fgb', 'gpkg', 'shp', 'geojson'])
def test_package_preserves_universe_and_blank_values(inputs, tmp_path, fmt):
    items, source = inputs
    payload = Selecao(attributes=[i['id'] for i in items], format=fmt).model_dump()
    package = dados.export_layer(payload)
    with zipfile.ZipFile(io.BytesIO(package)) as z:
        assert z.testzip() is None
        z.extractall(tmp_path)
        names = set(z.namelist())
    base = dados.PREFIXO
    manifest = json.loads(
        (tmp_path / f'{base}_metadados.json').read_text(encoding='utf-8')
    )
    assert set(manifest['arquivos']) == names
    assert manifest['join']['feicoes_entrada'] == manifest['join']['feicoes_saida'] == 3
    assert manifest['join']['campos']['vazio']['sem_valor'] == 3
    assert manifest['join']['campos']['parcial']['sem_observacao'] == 1
    layer = gpd.read_file(tmp_path / f'{base}.{fmt}').set_index('CD_MUN').sort_index()
    assert len(layer) == len(source)
    assert layer.geometry.geom_equals(source.geometry).all()
    fields = {i['field']: i['export_field'] for i in manifest['attributes']}
    assert layer[fields['vazio']].isna().all()
    assert layer.loc['3500001', fields['parcial']] == 0
    assert layer.loc[['3500002', '3500003'], fields['parcial']].isna().all()
    for ext, sep in [('csv', ','), ('txt', '\t')]:
        table = pd.read_csv(
            tmp_path / f'{base}_atributos.{ext}',
            sep=sep,
            dtype={'CD_MUN': str},
        ).set_index('CD_MUN').sort_index()
        assert list(table.columns) == list(layer.drop(columns='geometry').columns)
        assert table[fields['vazio']].isna().all()
        assert list(table.index) == list(source.index)
        assert table.loc['3500001', fields['parcial']] == 0
    wb = load_workbook(tmp_path / f'{base}_atributos.xlsx', read_only=True)
    rows = list(wb.active.iter_rows(max_col=len(layer.columns), values_only=True))
    wb.close()
    assert len(rows) == 4
    assert {r[0] for r in rows[1:]} == set(source.index)
    index = rows[0].index(fields['vazio'])
    assert all(r[index] is None for r in rows[1:])
    with (tmp_path / f'{base}_dicionario.csv').open(encoding='utf-8-sig') as f:
        glossary = list(csv.DictReader(f))
    assert {r['campo_exportado'] for r in glossary} == (
        set(layer.columns) - {'geometry'} | {'CD_MUN'}
    )
    assert all(r['campo_bruto'] and r['alias'] for r in glossary)
    report = (tmp_path / f'{base}_relatorio_join.txt').read_text(encoding='utf-8')
    assert 'Feições de entrada: 3; feições de saída: 3' in report
    if fmt == 'shp':
        assert {f'{base}{e}' for e in ['.shp', '.shx', '.dbf', '.prj', '.cpg']} <= names
    if fmt == 'gpkg':
        from osgeo import gdal
        assert manifest['alias_no_arquivo'] is True
        dataset = gdal.OpenEx(str(tmp_path / f'{base}.gpkg'), gdal.OF_VECTOR)
        definition = dataset.GetLayer(0).GetLayerDefn()
        aliases = {definition.GetFieldDefn(i).GetName(): definition.GetFieldDefn(i).GetAlternativeName()
                   for i in range(definition.GetFieldCount())}
        assert aliases['CD_MUN'] == dados.CAMPOS_FIXOS['CD_MUN'][0]
        assert all(aliases[i['export_field']] == i['alias'] for i in manifest['attributes'])
        dataset = None


@pytest.mark.parametrize('fmt', ['fgb', 'gpkg', 'shp', 'geojson'])
def test_geracao_registra_camada_pacote_e_metadados_no_storage(inputs, monkeypatch, request, fmt):
    from api.services import municipal_layer as service, ciclo_vida_arquivos as ciclo, saidas_storage
    from api.repositories import camada_geoespacial_repository as camadas, saidas_geoespaciais_repository as catalogo
    from api.services.arquivos_geoespaciais_memoria import vetor, ler_vetor
    items, source = inputs
    files = request.getfixturevalue('storage_api')
    malha = source.iloc[[i % len(source) for i in range(645)]].copy()
    malha.index = pd.Index([str(3500001 + i) for i in range(645)], name='CD_MUN')
    monkeypatch.setattr(dados, 'malha', malha.copy)
    monkeypatch.setattr(service, 'categoria', lambda code: {'id': code, 'nome': 'Social'})
    execution = str(uuid4())
    monkeypatch.setattr(ciclo, 'iniciar', lambda *a: execution)
    ended = []
    monkeypatch.setattr(ciclo, 'finalizar', lambda *a, **kw: ended.append((a, kw)))
    documents = []
    monkeypatch.setattr(catalogo, 'registrar_documento', lambda *a, **kw: documents.append(a[2]))

    def salvar_vetor(*, gdf, metadados, **kwargs):
        assert ciclo.execucao_atual.get() == execution
        content = vetor(gdf, 'resultado.gpkg', 'GPKG', layer='resultado')['resultado.gpkg']
        result = saidas_storage.enviar_bytes(execution, 'camadas', 'resultado.gpkg', content)
        metadados['caminho_arquivo'] = result['caminho']

    monkeypatch.setattr(camadas, 'salvar_vetor', salvar_vetor)
    payload = {'attributes': [i['id'] for i in items], 'format': fmt}
    package, result = service.gerar('social', payload, 'Camada de teste', SimpleNamespace(id=uuid4()))
    assert result['arquivo'].startswith(f'saidas-geoespaciais/execucoes/{execution}/camadas/')
    assert saidas_storage.conferir(result['pacote_caminho'], result['pacote_sha256']) == package
    actual = ler_vetor({'resultado.gpkg': files[result['arquivo']]}, 'resultado.gpkg')
    assert len(actual) == 645 and actual.crs.to_epsg() == 4674
    assert actual.CD_MUN.is_unique
    assert len(documents) == 2
    assert all(document['caminho'] in files for document in documents)
    assert ended == [((execution,), {})]
    assert not service._lock.locked()
