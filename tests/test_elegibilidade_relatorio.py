from tests.territorial_ogr_helpers import calcular_ogr
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from types import SimpleNamespace
from zipfile import ZipFile

import geopandas as gpd
import pytest
from pypdf import PdfReader
from shapely.geometry import box, Polygon, LineString
from fastapi import FastAPI
from fastapi.testclient import TestClient

from api.services import elegibilidade_relatorio as service
from api.services.arquivos_geoespaciais_memoria import vetor
from tests.territorial_ogr_helpers import serializar_fixture as serializar


@pytest.fixture
def analise(monkeypatch):
    fontes = [{'id': f'base-{i}', 'nome': f'base_tecnica_{i}', 'papel': 'restricao' if i < 4 else 'risco',
               'campo': f'campo_{i}', 'unidade': 'm²'} for i in range(7)]
    geometria = [box(5000000 + i*20, 7000000, 5000010 + i*20, 7000010) for i in range(63)]
    frame = gpd.GeoDataFrame({'ID_FEICAO': [f'Feição {i}' for i in range(63)],
                             'nome': [f'Demanda {i}' for i in range(63)]}, geometry=geometria, crs=5880)
    for i in range(7):
        frame[f'campo_{i}'] = [None, 0., 0.0000001] + [float(i+1)]*60
    frame['area_total_m2'] = 100.
    feicoes_base = [{'id': f'feicao-{i}', 'base_id': f'base-{i}', 'nome': f'trecho_do_setor_{i}',
                     'tipo': 'feicao', 'papel': fontes[i]['papel'], 'posicao': i, 'campo_nome': 'nome'} for i in range(7)]
    recortes = [{'demanda_pos': pos, 'base_id': f'base-{i}', 'feicao_id': f'feicao-{i}',
                 'area_m2': frame.iloc[pos][f'campo_{i}'],
                 'geometry': LineString([(5000000+pos*20,7000000),(5000000+pos*20,7000010)]) if pos == 1
                 else box(5000000+pos*20,7000000,5000000+pos*20+frame.iloc[pos][f'campo_{i}']/10,7000010)}
                for pos in range(1,63) for i in range(7)]
    nome = 'entrada_tecnica_elegibilidade_territorial.gpkg'
    dados = {}
    resultados = {}
    for i in (1, 2):
        auditoria = {'entrada': {'id': f'entrada-{i}', 'nome': f'entrada_tecnica_{i}'},
                     'campo_id': 'ID_FEICAO', 'feicoes': len(frame), 'campos': fontes,
                     'campo_area_total': 'area_total_m2', 'feicoes_base': feicoes_base,
                     'versao_saida': 2, 'camada_intersecoes': 'intersecoes'}
        conteudo = serializar(frame, auditoria, recortes, nome)[nome]
        arquivo = f'saidas-geoespaciais/saida-{i}/{nome}'
        dados[arquivo] = conteudo
        resultados[f'job-{i}'] = {'arquivo': arquivo, 'sha256': sha256(conteudo).hexdigest(),
            'lote_id': 'lote', 'total_entradas': 2, 'concluido_em': '2026-10-19T15:00:00+00:00',
            'elegibilidade': auditoria}
    consultar_persistido = service.jobs.consultar
    def consultar(ident, dono):
        if ident in resultados:
            if dono != 'autor':
                raise FileNotFoundError('Execução não encontrada.')
            return {'status': 'concluido', 'resultado': deepcopy(resultados[ident])}
        return consultar_persistido(ident, dono)
    monkeypatch.setattr(service.jobs, '_jobs', {})
    monkeypatch.setattr(service.jobs, 'consultar', consultar)
    def baixar(caminho):
        if caminho not in dados:
            raise FileNotFoundError(caminho)
        return dados[caminho]
    monkeypatch.setattr(service.storage_remoto, 'baixar', baixar)
    monkeypatch.setattr(service.storage_remoto, 'enviar_bytes', dados.__setitem__)
    monkeypatch.setattr(service.storage_remoto, 'mover', lambda a, b: dados.__setitem__(b, dados.pop(a)))
    nomes = {f'base-{i}': f'Bacia Hidrográfica do Setor {i}' for i in range(7)}
    nomes.update({f'feicao-{i}': f'Trecho do Setor {i}' for i in range(7)})
    nomes.update({'entrada-1': 'Demandas do Setor Norte', 'entrada-2': 'Demandas do Setor Sul'})
    return resultados, dados, nomes


@pytest.mark.parametrize(('original', 'esperado'), [
    ('áreas de preservação permanente (APP)', 'Áreas de Preservação Permanente (APP)'),
    ('ÁREAS DE RISCO', 'Áreas de Risco'), ('rede do DER em SP', 'Rede do DER em SP'),
    ('Áreas da XYZ', 'Áreas da XYZ'), ('bacia à jusante', 'Bacia à Jusante'),
])
def test_title_case_portugues_e_siglas(original, esperado):
    assert service.titulo(original) == esperado


def test_fontes_preservam_taxonomia_gravada_sem_reconsultar_dominio():
    taxonomia = {'tipo_codigo': 'inundacao', 'tipo_nome': 'Inundação',
                 'familia_codigo': 'perigo_fisico_geotecnico', 'familia_nome': 'Perigos físicos e geotécnicos'}
    fontes = service._fontes([{'elegibilidade': {
        'entrada': {'id': 'entrada', 'nome': 'Demandas'},
        'campos': [{'id': 'base', 'nome': 'Base', 'papel': 'restricao', 'adversidade': taxonomia}],
        'feicoes_base': [],
    }}])
    assert fontes[1]['adversidade'] == taxonomia
    assert fontes[1]['papel'] == 'restricao'
    assert 'adversidade' not in fontes[0]


def test_superficie_contexto_vinculo_autoria_e_extensao(analise, monkeypatch):
    fonte = gpd.GeoDataFrame({'nome':['Entorno','Distante']},
        geometry=[box(-48,-24,-46,-22),box(-40,-10,-39,-9)], crs=4326)
    original = fonte.copy()
    carregadas = []
    def carregar(base):
        carregadas.append(base)
        from tests.territorial_ogr_helpers import camada
        return camada(fonte), None, None
    monkeypatch.setattr(service.extracao_preparacao, 'carregar', carregar)
    bounds = [-47.5,-23.5,-46.5,-22.5]
    dados = service.superficie_elegibilidade('job-1','base-0','autor',bounds,.001)
    assert dados['papel'] == 'restricao'
    assert len(dados['geometrias']['features']) == 1
    assert dados['geometrias']['features'][0]['properties']['nome'] == 'Entorno'
    from shapely.geometry import shape
    assert shape(dados['geometrias']['features'][0]['geometry']).bounds == tuple(bounds)
    assert fonte.equals(original)
    assert 'Fonte original atual' in dados['origem']
    with pytest.raises(FileNotFoundError):
        service.superficie_elegibilidade('job-1','base-0','outro',bounds)
    with pytest.raises(ValueError, match='não pertence'):
        service.superficie_elegibilidade('job-1','base-nao-vinculada','autor',bounds)
    with pytest.raises(ValueError, match='inválida'):
        service.superficie_elegibilidade('job-1','base-0','autor',[0,0,float('nan'),1])
    assert carregadas == ['base-0']


def test_atributos_da_feicao_usam_alias_amigavel_e_recusam_vinculo_ou_fonte_divergente(analise, monkeypatch):
    fonte = gpd.GeoDataFrame({'nome': [f'trecho_do_setor_{i}' for i in range(7)], 'cd_uc': list(range(7)),
                              'area_ha': [1.5] * 7},
                             geometry=[box(-48, -24, -46, -22)] * 7, crs=4326)
    chamadas = []
    def carregar(base):
        chamadas.append(base)
        from tests.territorial_ogr_helpers import camada
        return camada(fonte), None, None
    monkeypatch.setattr(service.extracao_preparacao, 'carregar', carregar)
    dados = service.atributos_feicao('job-1', 'base-2', 2, 'autor')
    assert {c['campo']: c['alias'] for c in dados['campos']} == {'nome': 'Nome', 'cd_uc': 'Cd uc', 'area_ha': 'Area ha'}
    assert {c['campo']: c['valor'] for c in dados['campos']}['nome'] == 'trecho_do_setor_2'
    assert 'Fonte original atual' in dados['origem'] and chamadas == ['base-2']
    with pytest.raises(FileNotFoundError):
        service.atributos_feicao('job-1', 'base-2', 2, 'outro')
    with pytest.raises(ValueError, match='não pertence'):
        service.atributos_feicao('job-1', 'base-2', 5, 'autor')
    with pytest.raises(ValueError, match='não pertence'):
        service.atributos_feicao('job-1', 'base-nao-vinculada', 2, 'autor')
    fonte.loc[2, 'nome'] = 'Outro nome'
    with pytest.raises(ValueError, match='diverge'):
        service.atributos_feicao('job-1', 'base-2', 2, 'autor')


@pytest.mark.parametrize('hierarquizacao', [None, 'HIER-TESTE'])
def test_vinculo_cadastral_usa_codigo_original_e_nao_id_da_feicao(hierarquizacao):
    frame = gpd.GeoDataFrame({'ID_FEICAO':['feicao-1'], 'codigo':['D-001'], 'nome':['Nome da análise']},
                             geometry=[box(0, 0, 1, 1)], crs=5880)
    auditoria = {'campo_id':'ID_FEICAO', 'campos':[], 'hierarquizacao':hierarquizacao}
    linha = service._linhas(frame, auditoria)[0]
    assert linha['id'] == 'feicao-1'
    assert linha['codigo_demanda'] == ('D-001' if hierarquizacao else None)
    sem_codigo = service._linhas(frame.drop(columns='codigo'), auditoria)[0]
    assert sem_codigo['codigo_demanda'] is None


@pytest.mark.parametrize('nome', ['', 'risco_inundacao', '<script>', 'C:\\arquivo', 'Camada.gpkg'])
def test_alias_bruto_ou_invalido_e_recusado(nome):
    with pytest.raises(ValueError):
        service.titulo(nome)


def test_fontes_dinamicas_revisao_obrigatoria_e_aliases_exatos(analise):
    _, _, nomes = analise
    preparacao = service.preparar(['job-1', 'job-2'], 'autor')
    assert len(preparacao['fontes']) == 16
    assert all('_' not in f['nome'] for f in preparacao['fontes'])
    assert preparacao['previstas'] == preparacao['concluidas'] == 2
    with pytest.raises(ValueError, match='Revise'):
        service.montar(['job-1', 'job-2'], nomes, 'autor', False)
    with pytest.raises(ValueError, match='exatamente'):
        service.montar(['job-1', 'job-2'], {**nomes, 'extra': 'Extra'}, 'autor', True)
    with pytest.raises(ValueError, match='distintos'):
        service.montar(['job-1', 'job-2'], {**nomes, 'base-1': nomes['base-0']}, 'autor', True)


def test_tabela_mapa_e_pdf_completos_sem_nomes_tecnicos(analise):
    _, _, nomes = analise
    modelo, arquivos = service.montar(['job-1', 'job-2'], nomes, 'autor', True)
    assert len(arquivos) == 2 and len({n for n, _ in arquivos}) == 2
    entrada = modelo['entradas'][0]
    assert entrada['feicoes'] == 63 and len(entrada['campos']) == 7
    assert entrada['com_intersecao'] == 62 and entrada['sem_intersecao'] == 1
    assert entrada['linhas'][0]['areas_texto'] == ['—']*7
    assert entrada['linhas'][1]['areas_texto'] == ['0,0']*7
    assert entrada['linhas'][2]['areas_texto'] == ['0,0000001']*7
    assert len(entrada['mapa']['formas']) == 63 + 62*7
    assert not entrada['mapa']['formas'][0]['intersecao']
    assert entrada['mapa']['formas'][1]['intersecao']
    documento = PdfReader(BytesIO(service.pdf(modelo)))
    texto = '\n'.join(p.extract_text() for p in documento.pages)
    assert len(documento.pages) >= 7
    assert 'Feição 62' in texto and '0,0000000000001' in texto
    assert 'km²' in texto and ' ha' in texto and 'Trecho do Setor 6' in texto
    assert 'Bacia Hidrográfica do Setor 6' in texto
    assert 'Demandas do Setor Norte' in texto and 'Demandas do Setor Sul' in texto
    assert 'entrada_tecnica' not in texto and 'base_tecnica' not in texto and 'campo_0' not in texto
    assert 'Metodologia' in texto and 'Mapa-Síntese' in texto
    for pagina in documento.pages:
        assert float(pagina.mediabox.width) == pytest.approx(595.2756, abs=.01)
        assert float(pagina.mediabox.height) == pytest.approx(841.8898, abs=.01)


def test_pacote_persistido_preserva_geometrias_e_aliases(analise):
    _, dados, nomes = analise
    modelo, arquivos = service.montar(['job-1', 'job-2'], nomes, 'autor', True)
    zip_bytes, caminho = service.pacote(modelo, arquivos)
    assert dados[caminho] == zip_bytes and caminho.startswith('saidas-geoespaciais/')
    with ZipFile(BytesIO(zip_bytes)) as pacote:
        assert pacote.testzip() is None
        assert pacote.read('Relatório de Elegibilidade Territorial.pdf').startswith(b'%PDF')
        for nome, conteudo in arquivos:
            assert pacote.read(nome) == conteudo
        manifesto = json.loads(pacote.read('Metadados da Análise.json'))
        assert manifesto['nomes_revisados'] is True
        assert {f['nome'] for f in manifesto['fontes']} == set(nomes.values())
        assert manifesto['arquivos'][0]['sha256'] == sha256(arquivos[0][1]).hexdigest()


def test_pdf_e_pacote_informam_restricao_risco_ou_livre_por_demanda(analise):
    import csv as csv_lib
    from io import StringIO
    from openpyxl import load_workbook
    _, _, nomes = analise
    modelo, arquivos = service.montar(['job-1', 'job-2'], nomes, 'autor', True)
    texto = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(service.pdf(modelo))).pages)
    for trecho in ('Situação das Demandas quanto a Restrições e Riscos', 'A demanda é o objeto da análise',
                   'Distribuição das Demandas por Situação', 'Resumo por Demanda', 'Livre de restrições e riscos',
                   'Com restrição e risco', 'Tem 4 restrições', 'Anexo: Matrizes de Área por Camada'):
        assert trecho in texto, trecho
    assert 'Cobertura' not in texto
    assert texto.index('Resumo por Demanda') < texto.index('Detalhes por Demanda') < texto.index('Anexo: Matrizes')
    zip_bytes, _ = service.pacote(modelo, arquivos)
    with ZipFile(BytesIO(zip_bytes)) as pacote:
        conteudo = {'LEIA-ME.txt', 'Tabelas/Situação das Demandas.xlsx', 'Tabelas/Situação das Demandas.csv',
                    'Tabelas/Resumo da Elegibilidade.xlsx', 'Tabelas/Resumo da Elegibilidade.csv'}
        assert conteudo <= set(pacote.namelist())
        leia = pacote.read('LEIA-ME.txt').decode('utf-8-sig')
        assert 'A demanda é o objeto da análise' in leia and 'Livres de restrições e riscos: 2' in leia
        linhas = list(csv_lib.reader(StringIO(pacote.read('Tabelas/Situação das Demandas.csv').decode('utf-8-sig')),
                                     delimiter=';'))
        assert len(linhas) == 1 + 2 * (1 + 62 * 7)
        assert {l[3] for l in linhas[1:]} == {'Com restrição e risco', 'Livre de restrições e riscos'}
        livro = load_workbook(BytesIO(pacote.read('Tabelas/Situação das Demandas.xlsx')))
        assert livro.sheetnames == ['Situação das Demandas', 'Resumo da Elegibilidade', 'Restrições e Riscos',
                                    'Camadas Consultadas', 'Como Ler']
        resumo = list(csv_lib.reader(StringIO(pacote.read('Tabelas/Resumo da Elegibilidade.csv').decode('utf-8-sig')),
                                     delimiter=';'))
        assert resumo[0] == ['Conjunto de demandas', 'Objeto de demanda', 'Camadas de risco', 'Área de sobreposição (ha)',
                             'Porcentagem sobreposta', 'Índice de risco ambiental e territorial']
        assert len(resumo) == 1 + 2 * 1 and all(l[2] == '—' and l[5] == '1' for l in resumo[1:])
        assert 'Resumo da Elegibilidade' in texto and 'Índice de risco ambiental e territorial' in texto
        assert livro['Situação das Demandas'].max_row == 1 + 2 * 63
        assert livro['Restrições e Riscos'].max_row == 1 + 2 * (1 + 62 * 7)
        manifesto = json.loads(pacote.read('Metadados da Análise.json'))
        assert {d['nome'] for d in manifesto['documentos']} == conteudo


def test_resumo_da_elegibilidade_lista_demandas_sem_restricao_com_indice_de_zero_a_um():
    from api.services import elegibilidade_sintese as sintese
    item = lambda ident, base, papel, area: {'id': ident, 'base_id': base, 'papel': papel, 'area_m2': area,
                                              'nome': f'Item {ident}', 'percentual': area / 10, 'identificador': 1}
    entrada = {'nome': 'Conjunto', 'campos': [{'id': b, 'nome': f'Camada {b}'} for b in ('r1', 'r2', 'k')], 'linhas': [
        {'id': 'D1', 'nome': 'Com restrição', 'area_total_m2': 1000., 'sobreposicoes': [
            item('a', 'k', 'restricao', 100.), item('b', 'r1', 'risco', 50.)]},
        {'id': 'D2', 'nome': 'Risco em duas camadas', 'area_total_m2': 1000., 'risco_area_m2': 400., 'sobreposicoes': [
            item('c', 'r1', 'risco', 300.), item('d', 'r2', 'risco', 250.)]},
        {'id': 'D3', 'nome': 'Livre', 'area_total_m2': 1000., 'sobreposicoes': []},
        {'id': 'D4', 'nome': 'Toda sobreposta', 'area_total_m2': 1000., 'risco_area_m2': 1000., 'sobreposicoes': [
            item('e', 'r2', 'risco', 1000.)]},
        {'id': 'D5', 'nome': 'Só toca o limite', 'area_total_m2': 1000., 'sobreposicoes': [item('f', 'r1', 'risco', 0.)]}]}
    linhas = sintese.resumo_elegibilidade(entrada)
    assert [l['demanda'] for l in linhas] == ['Toda sobreposta', 'Risco em duas camadas', 'Livre', 'Só toca o limite']
    por_nome = {l['demanda']: l for l in linhas}
    duas = por_nome['Risco em duas camadas']
    assert duas['camadas'] == ['Camada r1', 'Camada r2'] and duas['area_m2'] == 400.
    assert duas['percentual'] == pytest.approx(40) and duas['indice'] == pytest.approx(.6)
    assert por_nome['Toda sobreposta']['indice'] == 0
    assert por_nome['Livre']['indice'] == 1 and por_nome['Livre']['camadas'] == [] and por_nome['Livre']['area_m2'] == 0
    assert por_nome['Só toca o limite']['indice'] == 1 and por_nome['Só toca o limite']['camadas'] == ['Camada r1']
    assert sintese.indice_risco_ambiental(5, 0) == 1 and sintese.indice_risco_ambiental(2000, 1000) == 0
    registro = sintese.resumo_para_registro({'entradas': [entrada]})
    assert registro['linhas'][1]['indice_de_risco_ambiental_e_territorial'] == pytest.approx(.6)
    assert registro['linhas'][1]['area_de_sobreposicao_ha'] == pytest.approx(.04)


def test_area_de_risco_une_camadas_sem_contar_duas_vezes_o_que_coincide():
    campos = {'r1': {'papel': 'risco'}, 'r2': {'papel': 'risco'}, 'k': {'papel': 'restricao'}}
    recortes = gpd.GeoDataFrame({'demanda_pos': [0, 0, 0, 1], 'base_id': ['r1', 'r2', 'k', 'r1']}, geometry=[
        box(5000000, 7000000, 5000006, 7000010), box(5000004, 7000000, 5000010, 7000010),
        box(5000000, 7000000, 5000010, 7000010), box(5000020, 7000000, 5000025, 7000010)], crs=5880)
    areas = service._area_uniao_risco(recortes, campos)
    assert areas[0] == pytest.approx(100.) and areas[1] == pytest.approx(50.)
    assert service._area_uniao_risco(recortes.iloc[0:0], campos) == {}
    assert service._area_uniao_risco(recortes[recortes['base_id'] == 'k'], campos) == {}


def test_situacao_da_demanda_e_linguagem_corrente_e_csv_neutraliza_formulas():
    from api.services import elegibilidade_sintese as sintese
    item = lambda ident, base, papel, area: {'id': ident, 'base_id': base, 'papel': papel, 'area_m2': area,
                                              'nome': f'Item {ident}', 'percentual': area / 10, 'identificador': 1}
    entrada = {'campos': [{'id': b, 'nome': f'Camada {b}'} for b in ('b1', 'b2', 'b3')], 'linhas': [
        {'id': 'D1', 'nome': 'Com ambos', 'area_total_m2': 1000., 'sobreposicoes': [
            item('a', 'b1', 'restricao', 100.), item('c', 'b2', 'restricao', 400.), item('d', 'b3', 'risco', 0.)]},
        {'id': 'D2', 'nome': 'Só risco', 'area_total_m2': 1000., 'sobreposicoes': [item('e', 'b3', 'risco', 5.)]},
        {'id': 'D3', 'nome': 'Só restrição', 'area_total_m2': 1000., 'sobreposicoes': [item('f', 'b1', 'restricao', 5.)]},
        {'id': 'D4', 'nome': 'Livre', 'area_total_m2': 1000., 'sobreposicoes': []}]}
    fichas = sintese.caracterizar(entrada)
    assert [f['situacao'] for f in fichas] == ['ambas', 'risco', 'restricao', 'livre']
    assert [f['nome'] for f in sintese.ordenar_por_situacao(fichas)] == ['Com ambos', 'Só restrição', 'Só risco', 'Livre']
    assert sintese.indicadores(fichas) == {'total': 4, 'ambas': 1, 'restricao': 1, 'risco': 1, 'livre': 1}
    assert sintese.descrever(fichas[0]) == 'Tem 2 restrições: Item a; Item c. Tem risco: Item d.'
    assert sintese.descrever(fichas[3]) == 'Livre de restrições e riscos nas camadas consultadas.'
    assert sintese._celula_csv('=1+1') == "'=1+1" and sintese._celula_csv('Normal') == 'Normal'
    assert sintese._celula_csv(0.5) == '0,5' and sintese._celula_csv(None) == ''


def test_nova_analise_reune_camadas_pdf_zip_e_revisoes_sem_sobrescrever(analise, monkeypatch):
    from api.services import hierarquizacao_service
    resultados, dados, nomes = analise
    pasta = service.pasta_saida('lote', 'HIER-001')
    for indice, resultado in enumerate(resultados.values(), 1):
        anterior = resultado['arquivo']
        resultado['arquivo'] = f'{pasta}/demanda_{indice}.gpkg'
        dados[resultado['arquivo']] = dados.pop(anterior)
        resultado['elegibilidade'].update(hierarquizacao='HIER-001', pasta_analise=pasta)
    monkeypatch.setattr(hierarquizacao_service, 'registrar_saida_geoespacial_fase1', lambda *args: None)
    monkeypatch.setattr(hierarquizacao_service, 'registrar_resultado_por_demanda', lambda *args: {'registrado': True})
    monkeypatch.setattr(hierarquizacao_service, 'atualizar_nomes_resultado', lambda *args: None)
    modelo, arquivos = service.montar(['job-1', 'job-2'], nomes, 'autor', True)
    service.salvar_nomes(['job-1', 'job-2'], modelo, 'autor')
    service.registrar_relatorios(['job-1', 'job-2'], modelo, 'autor')
    _, pacote = service.pacote(modelo, arquivos)
    pdf = modelo['relatorios_hierarquizacao'][0]['arquivo']
    assert modelo['pasta_analise'] == pasta
    assert pacote.rsplit('/', 1)[0] == pdf.rsplit('/', 1)[0] == pasta
    assert all(caminho.startswith(pasta + '/') for caminho in dados)
    assert all(service._caminho_revisao(r, 'autor').startswith(pasta + '/revisoes/') for r in resultados.values())
    anteriores = dict(dados)
    service.registrar_relatorios(['job-1', 'job-2'], modelo, 'autor')
    _, novo_pacote = service.pacote(modelo, arquivos)
    assert novo_pacote != pacote and modelo['relatorios_hierarquizacao'][0]['arquivo'] != pdf
    assert all(dados[caminho] == conteudo for caminho, conteudo in anteriores.items())


def test_saida_alterada_impede_resultados_e_exportacao(analise):
    resultados, dados, nomes = analise
    dados[resultados['job-1']['arquivo']] += b'alterado'
    with pytest.raises(ValueError, match='mudou'):
        service.montar(['job-1', 'job-2'], nomes, 'autor', True)


def test_saida_migrada_reutiliza_pasta_fisica_sem_regravar_geopackage(analise):
    resultados, dados, _ = analise
    resultado = resultados['job-1']
    arquivo_antigo = resultado['arquivo']
    pasta = service.pasta_saida('lote-migrado', 'HIER-001')
    resultado['arquivo'] = pasta + '/analise.gpkg'
    original = dados.pop(arquivo_antigo)
    dados[resultado['arquivo']] = original
    resultado['elegibilidade'].pop('pasta_analise', None)
    assert service._pasta_resultados([resultado]) == pasta
    assert service._caminho_revisao(resultado, 'autor').startswith(pasta + '/revisoes/')
    assert dados[resultado['arquivo']] == original


@pytest.mark.parametrize('segunda_hierarquizacao', [None, 'HIER-SUL'])
def test_pdf_automatico_associado_so_a_hierarquizacao_de_origem(analise, monkeypatch, segunda_hierarquizacao):
    from api.services import hierarquizacao_service
    resultados, dados, nomes = analise
    resultados['job-1']['elegibilidade']['hierarquizacao'] = 'HIER-NORTE'
    if segunda_hierarquizacao:
        resultados['job-2']['elegibilidade']['hierarquizacao'] = segunda_hierarquizacao
    chamadas = []
    monkeypatch.setattr(hierarquizacao_service, 'registrar_saida_geoespacial_fase1',
                        lambda *args: chamadas.append(args))
    monkeypatch.setattr(hierarquizacao_service, 'registrar_resultado_por_demanda', lambda *args: {'registrado': True})
    monkeypatch.setattr(hierarquizacao_service, 'atualizar_nomes_resultado', lambda *args: None)
    modelo, _ = service.montar(['job-1', 'job-2'], nomes, 'autor', True)
    service.registrar_relatorios(['job-1', 'job-2'], modelo, 'autor')
    assert len(chamadas) == (2 if segunda_hierarquizacao else 1)
    for i, (codigo, _, _, documento) in enumerate(chamadas):
        assert sha256(dados[documento['arquivo']]).hexdigest() == documento['sha256']
        texto = ' '.join(p.extract_text() for p in PdfReader(BytesIO(dados[documento['arquivo']])).pages)
        presente, ausente = ('Norte', 'Sul') if codigo == 'HIER-NORTE' else ('Sul', 'Norte')
        assert f'Demandas do Setor {presente}' in texto and f'Demandas do Setor {ausente}' not in texto
        assert modelo['relatorios_hierarquizacao'][i]['arquivo'] == documento['arquivo']
        planilha = documento['tabela_resumo']
        assert sha256(dados[planilha['arquivo']]).hexdigest() == planilha['sha256']
        assert planilha['arquivo'].endswith('.xlsx') and documento['resumo_elegibilidade']['titulo'] == 'Resumo da Elegibilidade'
    chamadas.clear()
    del resultados['job-1']['elegibilidade']['hierarquizacao']
    resultados['job-2']['elegibilidade'].pop('hierarquizacao', None)
    anteriores = dict(dados)
    service.registrar_relatorios(['job-1', 'job-2'], modelo, 'autor')
    assert not chamadas and dados == anteriores


def test_download_de_relatorio_exige_vinculo_autenticacao_e_assinatura(analise, monkeypatch):
    from api.deps.auth import require_authenticated
    from api.routers.hierarquizacoes import router
    from api.services import hierarquizacao_service
    arquivo = 'saidas-geoespaciais/relatorio.pdf'
    analise[1][arquivo] = b'%PDF-relatorio'
    registro = {'arquivo': analise[0]['job-1']['arquivo'], 'sha256': analise[0]['job-1']['sha256'],
                'relatorio': {'arquivo': arquivo, 'sha256': sha256(analise[1][arquivo]).hexdigest()}}
    monkeypatch.setattr(hierarquizacao_service.repo, 'get_by_codigo',
                        lambda codigo: {'relatorio_fase1': {'analises_geoespaciais': [registro]}})
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[require_authenticated] = lambda: SimpleNamespace(id='leitor')
    rota = '/ahp/hierarquizacoes/HIER-NORTE/relatorios/fase1/arquivo'
    with TestClient(app) as client:
        resposta = client.get(rota, params={'arquivo': arquivo})
        assert resposta.status_code == 200 and resposta.content == b'%PDF-relatorio'
        assert resposta.headers['content-type'] == 'application/pdf'
        assert 'attachment;' in resposta.headers['content-disposition']
        assert client.get(rota, params={'arquivo': 'saidas-geoespaciais/outro.pdf'}).status_code == 404
        assert client.get(rota.replace('/fase1/', '/fase2/'), params={'arquivo': arquivo}).status_code == 404
        analise[1][arquivo] += b'mudou'
        assert client.get(rota, params={'arquivo': arquivo}).status_code == 422
        from fastapi import HTTPException
        def sem_acesso():
            raise HTTPException(401, 'Autenticação necessária.')
        app.dependency_overrides[require_authenticated] = sem_acesso
        assert client.get(rota, params={'arquivo': arquivo}).status_code == 401


def test_reabrir_geopackage_reutiliza_saida_e_revisao_sem_recalcular(analise):
    resultados, dados, nomes = analise
    arquivo = resultados['job-1']['arquivo']
    reabertura = service.abrir(arquivo, 'autor')
    assert len(reabertura['execucoes']) == 1
    assert reabertura['revisados'] is False
    assert reabertura['concluidas'] == reabertura['previstas'] == 1
    nomes_unicos = {ident: nome for ident, nome in nomes.items() if ident != 'entrada-2'}
    modelo, _ = service.montar(reabertura['execucoes'], nomes_unicos, 'autor', True)
    service.salvar_nomes(reabertura['execucoes'], modelo, 'autor')
    reabertura = service.abrir(arquivo, 'autor')
    assert reabertura['revisados'] is True
    assert {f['id']: f['nome'] for f in reabertura['fontes']} == nomes_unicos
    assert dados[arquivo].startswith(b'SQLite format 3')
    dados[arquivo] += b'alterado'
    with pytest.raises(ValueError, match='mudou'):
        service.montar(reabertura['execucoes'], nomes_unicos, 'autor', True)


def test_reabrir_saida_de_outro_usuario_sem_historico_e_sem_recalcular(analise, monkeypatch):
    resultados, dados, nomes = analise
    arquivo = resultados['job-1']['arquivo']
    original = dados[arquivo]
    def nao_consultar_historico(*args, **kwargs):
        pytest.fail('A reabertura deve ler o GeoPackage, não procurar o histórico do autor.')
    monkeypatch.setattr(service.storage_remoto, 'listar', nao_consultar_historico)
    reabertura = service.abrir(arquivo, 'outro')
    nomes_unicos = {ident: nome for ident, nome in nomes.items() if ident != 'entrada-2'}
    modelo, arquivos = service.montar(reabertura['execucoes'], nomes_unicos, 'outro', True)
    assert len(modelo['entradas'][0]['linhas']) == 63
    assert modelo['entradas'][0]['linhas'][1]['areas'] == [0.] * 7
    assert modelo['entradas'][0]['linhas'][3]['areas'] == [float(i + 1) for i in range(7)]
    assert arquivos[0][1] == dados[arquivo] == original
    assert service.pdf(modelo).startswith(b'%PDF')
    monkeypatch.setattr(service.jobs, '_jobs', {})
    assert service.preparar(reabertura['execucoes'], 'outro')['fontes'] == reabertura['fontes']
    with pytest.raises(FileNotFoundError):
        service.preparar(reabertura['execucoes'], 'autor')


def test_reabrir_saida_compartilhada_nao_bloqueia_ou_sobrescreve_revisao_de_outro(analise):
    resultados, dados, nomes = analise
    arquivo = resultados['job-1']['arquivo']
    nomes_unicos = {ident: nome for ident, nome in nomes.items() if ident != 'entrada-2'}
    modelo, _ = service.montar(['job-1'], nomes_unicos, 'autor', True)
    service.salvar_nomes(['job-1'], modelo, 'autor')
    dados[arquivo + '.nomes.json'] = dados.pop(service._caminho_revisao(resultados['job-1'], 'autor'))
    revisao_autor = dados[arquivo + '.nomes.json']
    assert service.abrir(arquivo, 'autor')['revisados'] is True
    reabertura = service.abrir(arquivo, 'outro')
    assert reabertura['revisados'] is False
    modelo, _ = service.montar(reabertura['execucoes'], nomes_unicos, 'outro', True)
    service.salvar_nomes(reabertura['execucoes'], modelo, 'outro')
    assert service.abrir(arquivo, 'outro')['revisados'] is True
    assert dados[arquivo + '.nomes.json'] == revisao_autor
    assert service.abrir(arquivo, 'autor')['revisados'] is True


def test_rota_reabrir_geopackage_preserva_autorizacao_e_erros(analise, monkeypatch):
    from fastapi import HTTPException
    from api.deps.auth import require_geospatial_access
    from api.routers.extracao_atributos import router
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[require_geospatial_access] = lambda: SimpleNamespace(id='autor')
    arquivo = analise[0]['job-1']['arquivo']
    rota = '/extracao-atributos/elegibilidade/resultados/abrir'
    with TestClient(app) as client:
        resposta = client.post(rota, json={'arquivo': arquivo})
        assert resposta.status_code == 200
        assert len(resposta.json()['execucoes']) == 1
        assert client.post(rota, json={'arquivo': ''}).status_code == 422
        app.dependency_overrides[require_geospatial_access] = lambda: SimpleNamespace(id='outro')
        assert client.post(rota, json={'arquivo': arquivo}).status_code == 200
        assert client.post(rota, json={'arquivo': 'saidas-geoespaciais/inexistente.gpkg'}).status_code == 404
        analise[1][arquivo] = b'arquivo corrompido'
        assert client.post(rota, json={'arquivo': arquivo}).status_code == 422
        def sem_acesso():
            raise HTTPException(403, 'Acesso geoespacial não autorizado.')
        app.dependency_overrides[require_geospatial_access] = sem_acesso
        assert client.post(rota, json={'arquivo': arquivo}).status_code == 403


@pytest.mark.parametrize('arquivo', [
    '../saida.gpkg', '/saidas-geoespaciais/saida.gpkg', 'pasta-privada/saida.gpkg',
    'saidas-geoespaciais/../saida.gpkg', 'saidas-geoespaciais/saida.fgb',
    'saidas-geoespaciais\\saida.gpkg', 'saidas-geoespaciais//saida.gpkg',
])
def test_reabrir_recusa_caminho_ou_formato_invalido(analise, arquivo):
    with pytest.raises(ValueError, match='GeoPackage'):
        service.abrir(arquivo, 'outro')


def test_reabrir_recusa_geopackage_sem_metadados_e_nao_grava_referencia(analise):
    _, dados, _ = analise
    arquivo = 'saidas-geoespaciais/nao-elegibilidade.gpkg'
    frame = gpd.GeoDataFrame({'id': [1]}, geometry=[box(0, 0, 1, 1)], crs=5880)
    dados[arquivo] = vetor(frame, 'saida.gpkg', 'GPKG', layer='elegibilidade_territorial')['saida.gpkg']
    anteriores = dict(dados)
    with pytest.raises(ValueError, match='metadados'):
        service.abrir(arquivo, 'outro')
    assert dados == anteriores


@pytest.mark.parametrize(('campo', 'valor'), [
    ('entrada', None), ('campos', [{}]), ('campos', [{'id': 'base', 'nome': 'Base', 'campo': 'a', 'papel': []}]),
    ('feicoes_base', [{}]), ('campo_id', None), ('feicoes', 1), ('versao_saida', 1),
])
def test_reabrir_recusa_metadados_incompativeis_sem_registrar(analise, campo, valor):
    resultados, dados, _ = analise
    arquivo = resultados['job-1']['arquivo']
    auditoria = deepcopy(resultados['job-1']['elegibilidade'])
    auditoria[campo] = valor
    frame = gpd.read_file(BytesIO(dados[arquivo]), engine='pyogrio', layer='elegibilidade_territorial')
    dados[arquivo] = vetor(frame, 'saida.gpkg', 'GPKG', layer='elegibilidade_territorial',
                          layer_metadata={'SICARD_ELEGIBILIDADE': json.dumps(auditoria)})['saida.gpkg']
    anteriores = dict(dados)
    with pytest.raises(ValueError):
        service.abrir(arquivo, 'outro')
    assert dados == anteriores


def test_reabrir_nao_oculta_storage_indisponivel(analise, monkeypatch):
    def falhar(*args, **kwargs):
        raise RuntimeError('Storage indisponível.')
    monkeypatch.setattr(service.storage_remoto, 'baixar', falhar)
    with pytest.raises(RuntimeError, match='Storage indisponível'):
        service.abrir(analise[0]['job-1']['arquivo'], 'outro')


def test_nome_cadastral_nao_e_validado_como_alias_de_arquivo(analise, monkeypatch):
    _, _, nomes = analise
    ler = service._ler
    cadastral = 'PROPOSTA FEPASA - CAMELÓDROMO COMPLEXO DO EIXO FERROVIÁRIO - ASSIS/SP'
    longo = 'Demanda / setor _ especial ' * 10

    def com_nomes(resultado, *, analise='elegibilidade'):
        frame, dados = ler(resultado, analise=analise)
        frame.loc[0, 'nome'] = cadastral
        frame.loc[1, 'nome'] = longo
        return frame, dados

    monkeypatch.setattr(service, '_ler', com_nomes)
    modelo, _ = service.montar(['job-1', 'job-2'], nomes, 'autor', True)
    assert modelo['entradas'][0]['linhas'][0]['nome'] == cadastral
    assert modelo['entradas'][0]['linhas'][1]['nome'] == longo
    assert service.pdf(modelo).startswith(b'%PDF')
    with pytest.raises(ValueError, match='nome amigável'):
        service.montar(['job-1', 'job-2'], {**nomes, 'entrada-1': 'Arquivo/indevido'}, 'autor', True)


def test_parcial_identificado_e_lotes_distintos_recusados(analise):
    resultados, _, nomes = analise
    del nomes['entrada-2']
    modelo, _ = service.montar(['job-1'], nomes, 'autor', True)
    texto = ' '.join(p.extract_text() for p in PdfReader(BytesIO(service.pdf(modelo))).pages)
    assert 'RELATÓRIO PARCIAL' in texto
    resultados['job-2']['lote_id'] = 'outro'
    with pytest.raises(ValueError, match='mesma execução'):
        service.preparar(['job-1', 'job-2'], 'autor')
    with pytest.raises(FileNotFoundError):
        service.preparar(['job-1'], 'outro')


def test_painel_usa_geometrias_areas_e_feicoes_do_mesmo_arquivo(analise):
    _, _, nomes = analise
    modelo, _ = service.montar(['job-1','job-2'], nomes, 'autor', True)
    entrada = modelo['entradas'][0]
    assert len(entrada['geometrias']['features']) == 63
    assert len(entrada['intersecoes']['features']) == 62*7
    assert entrada['linhas'][0]['sobreposicoes'] == []
    contato = entrada['linhas'][1]['sobreposicoes'][0]
    assert contato['area_m2'] == contato['percentual'] == 0
    linha = entrada['linhas'][3]
    assert linha['area_total_m2'] == 100
    assert linha['perimetro_m'] == pytest.approx(40)
    assert linha['sobreposicoes'][-1]['nome'] == 'Trecho do Setor 6'
    assert linha['sobreposicoes'][-1]['area_m2'] == linha['areas'][-1] == 7.
    assert linha['sobreposicoes'][-1]['percentual'] == pytest.approx(7.)
    assert entrada['geometrias']['features'][3]['properties']['chave'] == linha['chave']


def test_perimetro_metrico_inclui_aneis_internos_e_reprojeta_crs(analise, monkeypatch):
    _, _, nomes = analise
    ler = service._ler
    def com_anel(resultado, *, analise='elegibilidade'):
        frame, dados = ler(resultado, analise=analise)
        frame.loc[0, frame.geometry.name] = Polygon(
            [(5000000,7000000),(5000010,7000000),(5000010,7000010),(5000000,7000010)],
            [[(5000002,7000002),(5000002,7000004),(5000004,7000004),(5000004,7000002)]])
        return frame.to_crs(4326), dados
    monkeypatch.setattr(service, '_ler', com_anel)
    modelo, _ = service.montar(['job-1','job-2'], nomes, 'autor', True)
    assert modelo['entradas'][0]['linhas'][0]['perimetro_m'] == pytest.approx(48, abs=1e-5)
    assert modelo['entradas'][0]['linhas'][1]['perimetro_m'] == pytest.approx(40, abs=1e-5)


def test_nome_sem_cadastro_persiste_na_revisao_e_no_pdf_regenerado(analise):
    resultados, dados, nomes = analise
    arquivo = resultados['job-1']['arquivo']
    original = dados[arquivo]
    resposta = service.renomear_demanda('job-1', '0', '  Demanda avulsa revisada  ', 'autor')
    assert resposta == {'execucao':'job-1', 'chave':'0', 'nome':'Demanda avulsa revisada'}
    service.renomear_demanda('job-1', '1', 'Outro nome / setor_especial', 'autor')
    modelo, arquivos = service.montar(['job-1','job-2'], nomes, 'autor', True)
    service.salvar_nomes(['job-1','job-2'], modelo, 'autor')
    reaberta = service.abrir(arquivo, 'autor')
    service.jobs._jobs.clear()
    regenerado, _ = service.montar(reaberta['execucoes'], {f['id']:f['nome'] for f in reaberta['fontes']}, 'autor', True)
    assert [l['nome'] for l in regenerado['entradas'][0]['linhas'][:2]] == [
        'Demanda avulsa revisada', 'Outro nome / setor_especial']
    assert modelo['entradas'][1]['linhas'][0]['nome'] == 'Demanda 0'
    assert dados[arquivo] == original == arquivos[0][1]
    pacote, _ = service.pacote(modelo, arquivos)
    with ZipFile(BytesIO(pacote)) as zipado:
        metadados = json.loads(zipado.read('Metadados da Análise.json'))
    assert metadados['entradas'][0]['nomes_demandas_revisados'] == {
        '0':'Demanda avulsa revisada', '1':'Outro nome / setor_especial'}
    texto = ' '.join(p.extract_text() for p in PdfReader(BytesIO(service.pdf(regenerado))).pages)
    assert 'Demanda avulsa revisada' in texto and 'Outro nome / setor_especial' in texto
    with pytest.raises(FileNotFoundError):
        service.renomear_demanda('job-1', '0', 'Não autorizado', 'outro')


@pytest.mark.parametrize(('chave','nome'), [('63','Nome'), ('-1','Nome'), ('01','Nome'), ('x','Nome'),
                                            ('0',' '), ('0','x'*201), ('0','Nome\ninválido')])
def test_revisao_demanda_recusa_chave_ou_nome_invalidos_sem_gravar(analise, chave, nome):
    _, dados, _ = analise
    anteriores = dict(dados)
    with pytest.raises(ValueError):
        service.renomear_demanda('job-1', chave, nome, 'autor')
    assert dados == anteriores


@pytest.mark.parametrize('campo', ['sha256', 'dono', 'nomes'])
def test_revisao_demanda_adulterada_e_explicita(analise, campo):
    resultados, dados, nomes = analise
    service.renomear_demanda('job-1', '0', 'Nome revisado', 'autor')
    caminho = service._caminho_nomes_demandas(resultados['job-1'], 'autor')
    revisao = json.loads(dados[caminho])
    revisao[campo] = {'999':'Inválido'} if campo == 'nomes' else 'inválido'
    dados[caminho] = json.dumps(revisao).encode()
    with pytest.raises(ValueError, match='revisão'):
        service.montar(['job-1','job-2'], nomes, 'autor', True)


def test_revisao_demanda_nao_oculta_falha_de_gravacao(analise, monkeypatch):
    monkeypatch.setattr(service.storage_remoto, 'enviar_bytes', lambda caminho, dados: None)
    with pytest.raises(RuntimeError, match='confirmar a gravação'):
        service.renomear_demanda('job-1', '0', 'Nome não salvo', 'autor')


def test_rota_revisao_de_demanda_exige_permissao_de_operacao(analise):
    from api.deps.auth import require_authenticated
    from api.routers.extracao_atributos import router
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[require_authenticated] = lambda: SimpleNamespace(id='autor', tipo_usuario='VISUALIZADOR')
    dados = analise[1]
    anteriores = dict(dados)
    with TestClient(app) as client:
        resposta = client.post('/extracao-atributos/elegibilidade/resultados/demanda/nome',
                               json={'execucao':'job-1', 'chave':'0', 'nome':'Sem permissão'})
    assert resposta.status_code == 403
    assert dados == anteriores


def test_nomes_confirmados_persistem_vinculados_ao_usuario_e_assinatura(analise):
    resultados, dados, nomes = analise
    assert service.preparar(['job-1','job-2'], 'autor')['revisados'] is False
    nomes['feicao-1'] = 'Feição Revisada'
    modelo, _ = service.montar(['job-1','job-2'], nomes, 'autor', True)
    service.salvar_nomes(['job-1','job-2'], modelo, 'autor')
    preparacao = service.preparar(['job-1','job-2'], 'autor')
    assert preparacao['revisados'] is True
    assert next(f for f in preparacao['fontes'] if f['id'] == 'feicao-1')['nome'] == 'Feição Revisada'
    caminho = service._caminho_revisao(resultados['job-1'], 'autor')
    revisao = json.loads(dados[caminho])
    revisao['sha256'] = 'assinatura diferente'
    dados[caminho] = json.dumps(revisao).encode()
    with pytest.raises(ValueError, match='proprietário ou à saída'):
        service.preparar(['job-1','job-2'], 'autor')


def test_nomes_de_feicoes_iguais_nao_unem_classes_e_legado_e_explicito(analise):
    resultados, _, nomes = analise
    nomes['feicao-1'] = nomes['feicao-0']
    modelo, _ = service.montar(['job-1','job-2'], nomes, 'autor', True)
    classes = modelo['entradas'][0]['linhas'][3]['sobreposicoes']
    assert classes[0]['nome'] == classes[1]['nome'] and classes[0]['id'] != classes[1]['id']
    resultados['job-1']['elegibilidade']['versao_saida'] = 1
    with pytest.raises(ValueError, match='Execute novamente'):
        service.montar(['job-1','job-2'], nomes, 'autor', True)


def test_mapa_preserva_aneis_internos():
    poligono = Polygon([(0,0),(10,0),(10,10),(0,10)], [[(2,2),(2,8),(8,8),(8,2)]])
    mapa = service._mapa(gpd.GeoDataFrame(geometry=[poligono], crs=5880), [True])
    assert len(mapa['formas'][0]['aneis']) == 2


def test_processamento_real_alimenta_painel_e_pdf_sem_recalcular(monkeypatch):
    from api.services import elegibilidade_processamento as processamento
    from api.services.controle_processamento import ControleProcessamento
    entrada = gpd.GeoDataFrame({'nome': ['Demanda Norte', 'Demanda Sul']},
        geometry=[box(5000000,7000000,5000010,7000010), box(5000020,7000000,5000030,7000010)], crs=5880)
    base = gpd.GeoDataFrame({'nome': ['Margens do Rio Claro', 'Reserva do Vale Verde']},
        geometry=[box(5000000,7000000,5000008,7000010), box(5000002,7000000,5000010,7000010)], crs=5880)
    recortes = []
    saida, audit = calcular_ogr(entrada, [('restricao','Proteção','base',base)],
                                          ControleProcessamento(), intersecoes=recortes)
    audit['entrada'] = {'id':'entrada','nome':'Demandas'}
    dados = processamento.serializar(saida, audit, recortes, 'saida.gpkg')['saida.gpkg']
    resultado = {'arquivo':'saidas-geoespaciais/saida.gpkg','sha256':sha256(dados).hexdigest(),'elegibilidade':audit}
    monkeypatch.setattr(service.jobs,'consultar',lambda ident,dono: {'status':'concluido','resultado':resultado})
    def baixar(caminho):
        if caminho != resultado['arquivo']:
            raise FileNotFoundError(caminho)
        return dados
    monkeypatch.setattr(service.storage_remoto,'baixar',baixar)
    nomes = {f['id']:f['nome'] for f in service._fontes([resultado])}
    modelo, arquivos = service.montar(['job'],nomes,'autor',True)
    linha = modelo['entradas'][0]['linhas'][0]
    assert linha['area_total_m2'] == 100 and linha['areas'] == [160.]
    assert [c['area_m2'] for c in linha['sobreposicoes']] == [80.,80.]
    assert sum(c['percentual'] for c in linha['sobreposicoes']) == 160.
    assert modelo['entradas'][0]['linhas'][1]['sobreposicoes'] == []
    assert arquivos[0][1] == dados
    texto = ' '.join(p.extract_text() for p in PdfReader(BytesIO(service.pdf(modelo))).pages)
    assert 'Margens do Rio Claro' in texto and 'Reserva do Vale Verde' in texto
    assert '80,00%' in texto and 'Demanda Sul' in texto
    recortes[0]['area_m2'] = 99.
    adulterado = processamento.serializar(saida,audit,recortes,'saida.gpkg')['saida.gpkg']
    resultado['sha256'] = sha256(adulterado).hexdigest()
    dados = adulterado
    with pytest.raises(ValueError, match='soma das interseções diverge'):
        service.montar(['job'],nomes,'autor',True)


def test_rotas_exportacao_autenticadas_e_revisadas(analise):
    from api.deps.auth import require_geospatial_access
    from api.routers.extracao_atributos import router
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[require_geospatial_access] = lambda: SimpleNamespace(id='autor')
    _, _, nomes = analise
    payload = {'execucoes': ['job-1','job-2'], 'nomes': nomes, 'revisados': True}
    rota = '/extracao-atributos/elegibilidade/resultados'
    with TestClient(app) as client:
        assert client.post(rota+'/preparar', json={'execucoes': payload['execucoes']}).status_code == 200
        assert client.post(rota+'/visualizar', json={**payload, 'revisados':False}).status_code == 422
        assert client.post(rota+'/visualizar', json=payload).status_code == 200
        resposta = client.post(rota+'/pacote', json=payload)
        assert resposta.status_code == 200 and resposta.headers['content-type'] == 'application/zip'
        assert 'filename*=' in resposta.headers['content-disposition']
        assert 'X-SICARD-Package' in resposta.headers
        assert client.post(rota+'/demanda/nome', json={'execucao':'job-1','chave':'0','nome':'Revisão pela API'}).json()['nome'] == 'Revisão pela API'
        assert client.post(rota+'/demanda/nome', json={'execucao':'job-1','chave':'0','nome':' '}).status_code == 422
        app.dependency_overrides[require_geospatial_access] = lambda: SimpleNamespace(id='outro')
        assert client.post(rota+'/pacote', json=payload).status_code == 404
        assert client.post(rota+'/demanda/nome', json={'execucao':'job-1','chave':'0','nome':'Sem acesso'}).status_code == 404
