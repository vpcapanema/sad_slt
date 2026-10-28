"""Contratos da preparação e do cálculo da superfície da Fase 2."""
import io
import json
from types import SimpleNamespace

import geopandas as gpd
import pandas as pd
import pytest
from openpyxl import load_workbook
from shapely.geometry import Point

from api.services import superficie_criterios as service


def test_reescalonamento_real_preserva_zero_e_ausentes():
    values = pd.Series([0., 10., 20., None])
    positive, audit = service.reescalonar(values, 'positiva')
    negative, _ = service.reescalonar(values, 'negativa')
    assert positive.iloc[:3].tolist() == [0, .5, 1]
    assert negative.iloc[:3].tolist() == [1, .5, 0]
    assert pd.isna(positive.iloc[3])
    assert audit == {'metodo': 'linear_min_max', 'minimo': 0, 'maximo': 20,
                     'relacao': 'positiva', 'ausentes': 1}


@pytest.mark.parametrize('values', [[2,2,2], [None,None], [0,float('inf')]])
def test_nao_inventa_reescalonamento(values):
    with pytest.raises(ValueError):
        service.reescalonar(pd.Series(values), 'positiva')


def test_cobertura_nao_e_contagem_numerica():
    basic = {'id':'a', 'coverage':0, 'detail':json.dumps({
        'registros_fonte_sp':103319, 'setores_sem_linha':0})}
    partial = {'id':'b', 'coverage':103319, 'detail':json.dumps({
        'registros_fonte_sp':100928, 'setores_sem_linha':2391})}
    combined = {'id':'c', 'detail':json.dumps({'atributos_origem':['a']})}
    catalog = {i['id']:i for i in [basic, partial, combined]}
    assert service.cobertura_comprovada(basic, 'setores_2022', catalog)
    assert not service.cobertura_comprovada(partial, 'setores_2022', catalog)
    assert service.cobertura_comprovada(combined, 'setores_2022', catalog)


def fixture_frame():
    return gpd.GeoDataFrame({'CD_MUN':['3500105','3500204','3500303'],
                            'NM_MUN':['A','B','C'], 'pop':[0,10,20], 'renda':[30,20,10]},
                           geometry=[Point(0,0),Point(1,1),Point(2,2)], crs=4674)


def test_camada_contem_cadastro_conceitos_e_pares_por_criterio():
    frame = fixture_frame()
    items = [{'id':'a','field':'pop','label':'População residente','source':'IBGE','year':2022,'theme':'populacao'},
             {'id':'b','field':'renda','label':'Renda média','source':'IBGE','year':2022,'theme':'renda'}]
    configs = [{'atributo_id':'a','dimensao':'Social','relacao':'positiva'},
               {'atributo_id':'b','dimensao':'Econômica','relacao':'negativa'}]
    output, criteria = service.preparar_frame(SimpleNamespace(layer=lambda _:frame), items, configs)
    assert output.columns[:5].tolist() == ['CD_MUN','NM_MUN','premissa','criterio','dimensao']
    assert output.geometry.equals(frame.geometry)
    for c in criteria:
        assert c['criterio'] == c['variavel'] + ' (favorabilidade territorial)'
        assert c['fenomeno_avaliado'] == 'favorabilidade territorial'
        assert c['premissa'].startswith('Quanto maior o valor de ')
        assert json.loads(output['dimensao'].iloc[0])[c['id']] == c['dimensao']
        assert c['campo_original'] in output and c['campo_reescalonado'] in output
    assert output[criteria[0]['campo_original']].tolist() == [0,10,20]
    assert service.compor(output, criteria, [.25,.75]).tolist() == [0,.5,1]


def test_composicao_nao_redistribui_pesos_nem_converte_nodata_em_zero():
    frame = pd.DataFrame({'a':[0,1,None],'b':[1,0,1]})
    criteria = [{'campo_reescalonado':'a'},{'campo_reescalonado':'b'}]
    result = service.compor(frame,criteria,[.2,.8])
    assert result.iloc[:2].tolist() == [.8,.2]
    assert pd.isna(result.iloc[2])
    with pytest.raises(ValueError,match='somar 1'):
        service.compor(frame,criteria,[1,1])


def test_matriz_individual_valida_consistencia_e_reciprocidade():
    result = service.pesos_matriz([[1,2],[.5,1]],2)
    assert result['weights'] == pytest.approx([2/3,1/3])
    with pytest.raises(ValueError,match='recíproca'):
        service.pesos_matriz([[1,2],[1,1]],2)
    with pytest.raises(ValueError,match='consistência'):
        service.pesos_matriz([[1,9,1/9],[1/9,1,9],[9,1/9,1]],3)


@pytest.mark.parametrize("origem", [{"matriz": [[1, 1], [1, 1]]}, {"ambiente_id": "historico"}])
def test_fase2_nao_calcula_novos_pesos_por_julgamento(monkeypatch, origem):
    monkeypatch.setattr(service, "obter", lambda *_: pytest.fail("Deve recusar antes de carregar a superfície"))
    with pytest.raises(ValueError, match="exclusivos da Fase 3"):
        service.calcular("superficie", {"operador": "ahp", **origem}, None)


def test_excel_usa_somente_criterios_da_superficie():
    row = {'id':'superficie','hierarquizacao_id':'hierarquia',
           'criterios':[{'id':'c001','dimensao':'Social','criterio':'Atendimento',
                         'premissa':'Quanto maior...', 'variavel':'População','relacao':'positiva',
                         'campo_reescalonado':'pop_r'}]}
    book = load_workbook(io.BytesIO(service.matriz_excel(row)), data_only=True)
    assert book['Comparação pareada'].cell(1,2).value == 'Atendimento'
    assert book['Comparação pareada'].cell(2,2).value == 1
    assert book['Vínculo'].cell(1,2).value == 'superficie'
    assert book['Dimensões de Critérios'].max_row == 13
    assert book['Matriz Crit Premissas v3'].cell(2,2).value == 'Atendimento'
    book.close()


def test_api_exige_autenticacao_e_renderiza_fluxo():
    from fastapi.testclient import TestClient
    from api.server import app
    from api.deps.auth import require_authenticated
    from api.services.session_service import SessionUser
    client = TestClient(app)
    assert client.get('/api/ahp/superficies-criterios/catalogo').status_code == 401
    user = SessionUser(id='12345678-1234-1234-1234-123456789012',email='a@example.test',
                       username='teste_ANALISTA',nome='Teste',tipo_usuario='ANALISTA')
    app.dependency_overrides[require_authenticated] = lambda:user
    try:
        response = client.get('/restrict/geoespacial/gerador-superficies-favorabilidade/')
        assert response.status_code == 200
        for ident in ('sc-viewer','sc-have','sc-mode','sc-judgment-new','sc-save'):
            assert f'id="{ident}"' in response.text
        phase = client.get('/restrict/hierarquizacao/fase-2/')
        assert phase.status_code == 200
        assert 'id="sc-app"' not in phase.text
    finally:
        app.dependency_overrides.pop(require_authenticated,None)


def test_api_julgamento_invalido_retorna_validacao():
    from fastapi import HTTPException
    from api.exceptions import DemandaValidationError
    from api.routers.superficie_criterios import invoke

    def invalid():
        raise DemandaValidationError('Julgamento não consolidado.')

    with pytest.raises(HTTPException) as exc:
        invoke(invalid)
    assert exc.value.status_code == 422
    assert exc.value.detail == 'Julgamento não consolidado.'
