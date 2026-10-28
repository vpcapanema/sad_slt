"""Média por área real de célula, auditoria persistida e compatibilidade vetorial."""
from io import BytesIO

import geopandas as gpd
import numpy as np
import pytest
from osgeo import gdal
from shapely.geometry import box

from api.services import favorabilidade_raster as raster
from api.services import elegibilidade_processamento as processo, elegibilidade_relatorio as relatorio
from api.services import agregacao_favorabilidade as agregacao, favorabilidade_exportacao as exportacao
from api.services.controle_processamento import ControleProcessamento, ProcessamentoCancelado
from tests.territorial_ogr_helpers import camada


@pytest.fixture
def fonte(tmp_path, monkeypatch):
    caminho = tmp_path / 'indice.tif'
    ds = gdal.GetDriverByName('GTiff').Create(str(caminho), 2, 2, 1, gdal.GDT_Float64)
    ds.SetGeoTransform((0, 10, 0, 20, 0, -10))
    ds.SetProjection(processo.ogr.referencia(5880).ExportToWkt())
    ds.GetRasterBand(1).SetNoDataValue(-9999)
    ds.GetRasterBand(1).WriteArray(np.array([[.2, .8], [0, -9999]]))
    ds.Close()
    monkeypatch.setattr(raster.storage, 'resolver', lambda _: caminho)
    ident = 'storage:superficies-indices/indice.tif'
    meta = raster.preparar(ident)
    return raster.SuperficieRaster(ident, meta['revisao_fonte']), caminho


def executar(fonte, geometrias=None, bases=None):
    entrada = camada(gpd.GeoDataFrame({'nome': ['Borda', 'Zero', 'Externa']},
        geometry=geometrias or [box(5, 10, 20, 20), box(0, 0, 20, 10), box(30, 0, 40, 10)], crs=5880))
    recortes = []
    saida, audit = processo.calcular(entrada, bases or [('rede', 'Rede', fonte.id, fonte)],
                                   ControleProcessamento(), analise='favorabilidade', intersecoes=recortes)
    audit['entrada'] = {'id': 'entrada', 'nome': 'Demandas'}
    return entrada, saida, audit, recortes


def test_media_ponderada_bordas_nodata_zero_e_saida_persistida(fonte):
    fonte, caminho = fonte
    original = caminho.read_bytes()
    entrada, saida, audit, recortes = executar(fonte)
    base = audit['campos'][0]
    assert [r[base['campo_indice']] for r in saida.registros] == [pytest.approx(.6), 0, None]
    assert [r[base['campo_cobertura']] for r in saida.registros] == [100, 50, 0]
    assert [r['area_m2'] for r in recortes] == [50, 100, 100]
    assert audit['indices_agregados'] is True
    assert base['raster']['banda'] == 1
    assert [g.ExportToWkb() for g in entrada.geometrias] == [g.ExportToWkb() for g in saida.geometrias]
    dados = processo.serializar(saida, audit, recortes, 'saida.gpkg', analise='favorabilidade')['saida.gpkg']
    persistido = relatorio._auditoria(dados, analise='favorabilidade')
    assert persistido['campos'] == audit['campos']
    frame = gpd.read_file(BytesIO(dados), layer='favorabilidade_territorial')
    linhas = relatorio._linhas(frame, audit)
    nomes = {f['id']: f['nome'] for f in [audit['entrada'], *audit['campos'], *audit['feicoes_base']]}
    relatorio._painel(frame, dados, audit, nomes, linhas, somente_atributos=True, analise='favorabilidade')
    assert linhas[0]['rasters'][0]['indice'] == pytest.approx(.6)
    assert linhas[0]['rasters'][0]['indice'] != pytest.approx(.5)  # Não é média simples.
    assert linhas[1]['rasters'][0] == {
        'base_id': fonte.id, 'papel': 'rede', 'indice': 0, 'area_valida_m2': 100,
        'cobertura_percentual': 50, 'area_sem_dados_m2': 100, 'celulas_validas': 1,
        'metrica': 'media_ponderada_area'}
    assert linhas[2]['rasters'][0]['indice'] is None
    for medida in ('media', 'maximo', 'cv', 'moda'):
        assert agregacao.valores_linha(linhas[0], {'medida': medida}) == {'grade': None, 'rede': pytest.approx(.6)}
        assert agregacao.valores_linha(linhas[1], {'medida': medida})['rede'] == 0
    modelo = {'entradas': [{'id': 'job', 'nome': 'Demandas', 'campos': audit['campos'], 'linhas': linhas}]}
    resumo = exportacao.sintese_raster(modelo)
    assert resumo[0][7:10] == [pytest.approx(.6), 150, 100]
    assert resumo[2][7:10] == [None, 0, 0]
    adulterado = frame.copy()
    adulterado.loc[0, base['campo_indice']] = .5
    with pytest.raises(ValueError, match='diverge'):
        relatorio._painel(adulterado, dados, audit, nomes, relatorio._linhas(adulterado, audit),
                         somente_atributos=True, analise='favorabilidade')
    assert caminho.read_bytes() == original
    from api.services.hierarquizacao_service import _fase2_da_demanda
    registrado = _fase2_da_demanda(linhas[0], {fonte.id: 'Rede'}, {}, {}, agregacao={'medida': 'maximo'})
    assert registrado['indice_favorabilidade_rede'] == pytest.approx(.6)
    assert registrado['extracao_raster'][0]['cobertura_percentual'] == 100
    assert registrado['agregacao']['unidade_por_dimensao']['rede'] == 'índice'
    assert agregacao.carregar([{'favorabilidade': audit}], 'autor') == {'medida': 'media', 'percentil': None}


def test_vetor_grade_e_raster_rede_independentes(fonte):
    fonte, _ = fonte
    grade = camada(gpd.GeoDataFrame({'fav_territorial_dim': [.3]}, geometry=[box(0, 0, 30, 30)], crs=5880))
    _, saida, audit, _ = executar(fonte, bases=[('grade', 'Grade', 'g', grade), ('rede', 'Rede', fonte.id, fonte)])
    assert saida.registros[0]['fg_grade'] == 150
    assert saida.registros[0][audit['campos'][1]['campo_indice']] == pytest.approx(.6)
    with pytest.raises(ValueError, match='única superfície raster'):
        executar(fonte, bases=[('rede', 'Grade', 'g', grade), ('rede', 'Rede', fonte.id, fonte)])


@pytest.mark.parametrize('valor', [-.01, 1.01, float('inf'), float('nan')])
def test_indice_invalido_nao_e_convertido_em_zero(fonte, valor):
    fonte, caminho = fonte
    ds = gdal.Open(str(caminho), gdal.GA_Update)
    ds.GetRasterBand(1).WriteArray(np.array([[valor, .8], [0, -9999]]))
    ds.Close()
    with pytest.raises(ValueError, match='Índice raster inválido'):
        executar(raster.SuperficieRaster(fonte.id))


def test_revisao_cancelamento_e_banda(fonte):
    fonte, caminho = fonte
    with pytest.raises(ValueError, match='mudou'):
        raster.preparar(fonte.id, 'versao-antiga')
    controle = ControleProcessamento()
    controle.cancelar()
    with pytest.raises(ProcessamentoCancelado):
        list(raster.celulas(fonte, camada(gpd.GeoDataFrame(geometry=[box(0, 0, 10, 10)], crs=5880)), controle))
    ds = gdal.GetDriverByName('GTiff').Create(str(caminho), 2, 2, 2, gdal.GDT_Float32)
    ds.SetProjection(processo.ogr.referencia(5880).ExportToWkt())
    ds.SetGeoTransform((0, 10, 0, 20, 0, -10))
    ds.Close()
    with pytest.raises(ValueError, match='única banda'):
        raster.preparar(fonte.id)


def test_aneis_internos_preservam_area_e_pesos(fonte):
    fonte, _ = fonte
    demanda = box(0, 10, 20, 20).difference(box(1, 11, 6, 19))
    entrada = camada(gpd.GeoDataFrame(geometry=[demanda], crs=5880))
    celulas = list(raster.celulas(fonte, entrada, ControleProcessamento()))
    assert [c[4] for c in celulas] == [60, 100]
    assert raster.resumir([(c[3], c[4]) for c in celulas], demanda.area)['indice'] == pytest.approx(.575)


def test_relatorio_mapa_e_planilhas_usam_recortes_historicos(fonte, monkeypatch):
    from hashlib import sha256
    from openpyxl import load_workbook
    from pypdf import PdfReader
    fonte, caminho = fonte
    from tests.relatorio_cartografia_helpers import preparar_fontes
    preparar_fontes(monkeypatch)
    ds = gdal.Open(str(caminho), gdal.GA_Update)
    ds.SetGeoTransform((5000000, 10, 0, 7000020, 0, -10))
    ds.Close()
    fonte = raster.SuperficieRaster(fonte.id)
    from shapely.affinity import translate
    geometrias = [translate(g, xoff=5000000, yoff=7000000)
                  for g in (box(5, 10, 20, 20), box(0, 0, 20, 10), box(30, 0, 40, 10))]
    _, saida, audit, recortes = executar(fonte, geometrias=geometrias)
    dados = processo.serializar(saida, audit, recortes, 'saida.gpkg', analise='favorabilidade')['saida.gpkg']
    resultado = {'execucao_id': 'job', 'arquivo': 'saidas-geoespaciais/saida.gpkg',
                 'sha256': sha256(dados).hexdigest(), 'favorabilidade': audit, 'total_entradas': 1}
    monkeypatch.setattr(relatorio, '_selecionar', lambda *args, **kwargs: [resultado])
    monkeypatch.setattr(relatorio.storage_remoto, 'baixar', lambda _: dados)
    nomes = {f['id']: f['nome'] for f in relatorio._fontes([resultado], analise='favorabilidade')}
    assert fonte.id in nomes
    modelo, arquivos = relatorio.montar(['job'], nomes, 'autor', True, analise='favorabilidade')
    assert modelo['entradas'][0]['campos'][0]['tipo'] == 'raster'
    assert modelo['entradas'][0]['linhas'][0]['rasters'][0]['indice'] == pytest.approx(.6)
    superficie = relatorio.superficie_favorabilidade('job', fonte.id, 'autor')
    assert superficie['tipo'] == 'raster'
    assert len(superficie['geometrias']['features']) == 3
    assert [f['properties']['favorabilidade']['indice'] for f in superficie['geometrias']['features']] == [.2, .8, 0]
    texto = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(relatorio.pdf(modelo))).pages)
    assert 'Média ponderada pela área' in texto and '0,60' in texto
    assert 'Cobertura válida: 50,00%' in texto
    documentos = dict(exportacao.documentos(modelo, 'Relatorio.pdf', [n for n, _ in arquivos]))
    assert b'0,6' in documentos['Tabelas/Síntese Raster.csv']
    livro = load_workbook(BytesIO(documentos['Tabelas/Favorabilidade Territorial.xlsx']))
    assert livro['Síntese Raster'].cell(2, 8).value == pytest.approx(.6)
    assert livro['Síntese Raster'].cell(3, 10).value == 50


def test_crs_geografico_transformado_sem_reamostrar_valores(tmp_path, monkeypatch):
    caminho = tmp_path / 'geografico.tif'
    ds = gdal.GetDriverByName('GTiff').Create(str(caminho), 1, 1, 1, gdal.GDT_Float64)
    ds.SetGeoTransform((-47, .001, 0, -23, 0, -.001))
    ds.SetProjection(processo.ogr.referencia(4326).ExportToWkt())
    ds.GetRasterBand(1).Fill(.75)
    ds.Close()
    monkeypatch.setattr(raster.storage, 'resolver', lambda _: caminho)
    demanda = camada(gpd.GeoDataFrame(geometry=[box(-46.9998, -23.0008, -46.9992, -23.0002)], crs=4326))
    metrica = processo.ogr.reproject(demanda, 5880)
    valores = list(raster.celulas(raster.SuperficieRaster('storage:superficies-indices/geografico.tif'),
                                metrica, ControleProcessamento()))
    resumo = raster.resumir([(c[3], c[4]) for c in valores], processo.ogr.area(metrica.geometrias[0]))
    assert resumo['indice'] == .75
    assert resumo['cobertura_percentual'] == pytest.approx(100)


def test_preparar_endpoint_e_revisao_fisica(fonte):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from types import SimpleNamespace
    from api.deps.auth import require_geospatial_access
    from api.routers.extracao_atributos import router
    fonte, caminho = fonte
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[require_geospatial_access] = lambda: SimpleNamespace(id='autor')
    with TestClient(app) as client:
        rota = '/extracao-atributos/favorabilidade/preparar-raster'
        resposta = client.post(rota, json={'id': fonte.id})
        assert resposta.status_code == 200
        assert resposta.json()['revisao_fonte'] == fonte.revisao
        assert client.post(rota, json={'id': fonte.id, 'revisao': 'antiga'}).status_code == 422
        assert client.post(rota, json={'id': fonte.id.replace('indice.tif', 'superficie_favorabilidade_rede_distancia_metros_200m.tif')}).status_code == 422
        app.dependency_overrides.clear()
        assert client.post(rota, json={'id': fonte.id}).status_code in (401, 403)
