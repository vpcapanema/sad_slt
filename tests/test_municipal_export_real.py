"""Reabre os quatro formatos usando cópia autorizada de malha e observações reais.

Snapshot JSON: catalog, chosen, malha (FeatureCollection) e rows (atributo_id,
municipios, valores). Nenhuma consulta ou gravação na produção durante o teste.
"""
import csv
import io
import json
import os
import pytest
import tempfile
import time
import zipfile
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

import geopandas as gpd
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from api.services import base_municipal as dados
from api.services import municipal_layer as service


@pytest.mark.skipif(not os.getenv('SICARD_MUNICIPAL_SNAPSHOT'), reason='Snapshot real autorizado não fornecido')
def test_pacotes_com_snapshot_real():
    snapshot=json.loads(Path(os.environ['SICARD_MUNICIPAL_SNAPSHOT']).read_text())
    source=gpd.GeoDataFrame.from_features(snapshot['malha']['features'],crs=4674).set_index('CD_MUN')
    items=snapshot['chosen']
    class Connection:
        def execute(self,*args): return self
        def fetchall(self): return snapshot['rows']
    @contextmanager
    def connection(): yield Connection()

    results=[]
    with patch.object(dados,'malha',lambda:source.copy()), patch.object(dados,'catalog',lambda:snapshot['catalog']), patch.object(dados,'get_connection',connection):
        expected=dados.layer(items).set_index('CD_MUN').sort_index()
        for fmt in ['fgb','gpkg','shp','geojson']:
            started=time.monotonic()
            with tempfile.TemporaryDirectory(prefix='sicard-real-'+fmt+'-') as temp:
                folder=Path(temp)
                package,path,manifest,frame=service.materializar({'attributes':[i['id'] for i in items],'format':fmt},folder)
                assert len(frame)==len(source)==645
                actual=frame.set_index('CD_MUN').sort_index()
                assert actual.geometry.geom_equals(expected.geometry).all()
                for item in manifest['attributes']:
                    np.testing.assert_allclose(actual[item['export_field']].to_numpy(dtype=float),expected[item['field']].to_numpy(dtype=float),equal_nan=True,rtol=1e-12,atol=1e-12)
                with zipfile.ZipFile(io.BytesIO(package)) as archive:
                    assert archive.testzip() is None
                    assert set(archive.namelist())==set(manifest['arquivos'])
                base=dados.PREFIXO
                if fmt=='shp': assert all((folder/(base+suffix)).exists() for suffix in ['.shp','.shx','.dbf','.prj','.cpg'])
                with (folder/f'{base}_dicionario.csv').open(encoding='utf-8-sig') as stream:
                    glossary=list(csv.DictReader(stream))
                assert len(glossary)==len(items)+4
                assert {r['campo_exportado'] for r in glossary}==set(frame.columns)-{'geometry'}
                assert all(r['alias'] and r['campo_bruto'] for r in glossary)
                numeric=['AREA_KM2']+[i['export_field'] for i in manifest['attributes']]
                for extension,sep in [('csv',','),('txt','\t')]:
                    table=pd.read_csv(folder/f'{base}_atributos.{extension}',sep=sep,dtype={'CD_MUN':str}).set_index('CD_MUN').sort_index()
                    assert list(table.index)==list(actual.index)
                    assert table.NM_MUN.tolist()==actual.NM_MUN.tolist()
                    np.testing.assert_allclose(table[numeric].to_numpy(dtype=float),actual[numeric].to_numpy(dtype=float),equal_nan=True,rtol=1e-12,atol=1e-12)
                wb=load_workbook(folder/f'{base}_atributos.xlsx',read_only=True)
                rows=list(wb.active.iter_rows(max_col=len(frame.columns)-1,values_only=True));wb.close()
                table=pd.DataFrame(rows[1:],columns=rows[0]).set_index('CD_MUN').sort_index()
                assert len(table)==645 and list(table.index)==list(actual.index)
                np.testing.assert_allclose(table[numeric].to_numpy(dtype=float),actual[numeric].to_numpy(dtype=float),equal_nan=True,rtol=1e-12,atol=1e-12)
                report=(folder/f'{base}_relatorio_join.txt').read_text()
                assert 'Feições de entrada: 645; feições de saída: 645' in report
                counts=manifest['join']['campos']
                assert any(c['sem_valor']==645 for c in counts.values())
                assert any(0<c['sem_valor']<645 for c in counts.values())
                result=dict(formato=fmt,feicoes=len(actual),indicadores=len(items),arquivos=len(manifest['arquivos']),bytes_zip=len(package),segundos=round(time.monotonic()-started,2),status='OK')
                results.append(result)
                print(json.dumps(result),flush=True)
