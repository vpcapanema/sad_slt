"""Modo enriquecimento: um registro por feição (ou trecho municipal) com regras por base."""
import geopandas as gpd
import pytest
from shapely.geometry import LineString, Point, Polygon, box

from api.services.extracao_atributos_enriquecimento import enriquecer, medida, preparar

X, Y = 5_000_000, 7_500_000


def gdf(geometrias, **campos):
    return gpd.GeoDataFrame(campos, geometry=geometrias, crs=5880)


def tema(*camadas, nome='Tema'):
    return [{'id': 'tema', 'nome': nome,
             'camadas': [{'id': c[0], 'nome': c[0], 'frame': c[1], 'regra': c[2] if len(c) > 2 else None}
                         for c in camadas]}]


MUNICIPIOS = gdf([box(X, Y, X + 100, Y + 100), box(X + 100, Y, X + 200, Y + 100)],
                 CD_MUN=['3500001', '3500002'], NM_MUN=['Alfa', 'Beta'])


def test_linha_cortada_por_municipio_leva_atributos_da_propria_unidade():
    # 250 m de linha: 100 em Alfa, 100 em Beta, 50 fora de qualquer município.
    entrada = gdf([LineString([(X, Y + 50), (X + 250, Y + 50)])], proj_id=['P1'])
    saida = enriquecer(entrada, tema(('municipios', MUNICIPIOS, {'papel': 'recorte', 'prefixo': 'mun'})))
    linhas = saida['camadas']['linhas'].to_crs(5880).sort_values('mun_fid_base', na_position='last')
    assert list(linhas['mun_NM_MUN'].fillna('fora')) == ['Alfa', 'Beta', 'fora']
    assert [round(medida(g, 1)) for g in linhas.geometry] == [100, 100, 50]
    assert list(linhas['proj_id']) == ['P1'] * 3 and set(linhas['fid_origem']) == {0}
    assert list(saida['camadas']['linhas']['id_registro']) == [1, 2, 3]
    assert saida['relatorio']['camadas']['linhas']['etapas'][0]['fora_das_unidades'] == 1


def test_maior_sobreposicao_escolhe_a_feicao_que_cobre_mais_o_trecho():
    entrada = gdf([LineString([(X, Y), (X + 100, Y)])], proj_id=['P1'])
    # Feição 0 cobre 20 m; feição 1 cobre 80 m.
    base = gdf([box(X - 10, Y - 5, X + 20, Y + 5), box(X + 20, Y - 5, X + 110, Y + 5)], nome=['pequena', 'grande'])
    maior = enriquecer(entrada, tema(('uc', base, {'prefixo': 'uc'})))['camadas']['linhas'].iloc[0]
    assert maior['uc_nome'] == 'grande' and maior['uc_n_feicoes'] == 2 and maior['uc_fid_base'] == 1
    assert round(maior['uc_comprimento_comum_m']) == 80
    primeira = enriquecer(entrada, tema(('uc', base, {'prefixo': 'uc', 'multiplicidade': 'primeira'})))
    assert primeira['camadas']['linhas'].iloc[0]['uc_nome'] == 'pequena'


def test_todas_duplica_e_resumo_agrega_sem_perder_registros():
    entrada = gdf([Point(X + 5, Y + 5), Point(X + 500, Y + 500)], proj_id=['dentro', 'fora'])
    base = gdf([box(X, Y, X + 10, Y + 10), box(X, Y, X + 20, Y + 20)], nome=['A', 'B'], area=[1.5, 2.5])
    todas = enriquecer(entrada, tema(('b', base, {'prefixo': 'b', 'multiplicidade': 'todas'})))['camadas']['pontos']
    assert list(todas['proj_id']) == ['dentro', 'dentro', 'fora']
    assert list(todas['b_nome'].fillna('')) == ['A', 'B', '']
    resumo = enriquecer(entrada, tema(('b', base, {'prefixo': 'b', 'multiplicidade': 'resumo'})))['camadas']['pontos']
    dentro, fora = resumo.iloc[0], resumo.iloc[1]
    assert dentro['b_nome'] == 'A | B' and dentro['b_area'] == 4.0 and dentro['b_n_feicoes'] == 2
    assert fora['b_n_feicoes'] == 0 and (fora['b_nome'] is None or fora['b_nome'] != fora['b_nome'])


def test_ligacao_por_atributo_com_chave_numerica_e_texto():
    entrada = gdf([LineString([(X + 10, Y + 50), (X + 190, Y + 50)])], proj_id=['P1'])
    indicadores = gdf([Point(0, 0), Point(1, 1)], CD_MUN=[3500001, 3500002], pib=[10.0, 20.0])
    categorias = [
        {'id': 'recorte', 'nome': 'Municípios', 'camadas': [{
            'id': 'm', 'nome': 'm', 'frame': MUNICIPIOS,
            'regra': {'papel': 'recorte', 'prefixo': 'mun', 'campos': ['CD_MUN']}}]},
        {'id': 'eco', 'nome': 'Econômico', 'camadas': [{'id': 'i', 'nome': 'i', 'frame': indicadores, 'regra': {
            'ligacao': 'atributo', 'chave_entrada': 'mun_CD_MUN', 'chave_base': 'CD_MUN', 'prefixo': 'eco',
            'campos': ['pib'], 'apelidos': {'pib': 'PIB municipal'}}}]},
    ]
    saida = enriquecer(entrada, categorias)
    linhas = saida['camadas']['linhas'].sort_values('mun_CD_MUN')
    assert list(linhas['eco_pib']) == [10.0, 20.0]
    entrada_dic = next(d for d in saida['dicionario'] if d['campo'] == 'eco_pib')
    assert entrada_dic['apelido'] == 'PIB municipal' and entrada_dic['tema'] == 'Econômico'
    assert 'mun_NM_MUN' not in linhas.columns, 'só os campos escolhidos'


def test_buffer_transforma_pontos_da_base_em_areas():
    entrada = gdf([Point(X + 30, Y), Point(X + 80, Y)], proj_id=['perto', 'longe'])
    contaminadas = gdf([Point(X, Y)], razao=['Posto'])
    sem = enriquecer(entrada, tema(('ac', contaminadas, {'prefixo': 'ac'})))['camadas']['pontos']
    assert list(sem['ac_n_feicoes']) == [0, 0]
    com = enriquecer(entrada, tema(('ac', contaminadas, {'prefixo': 'ac', 'preparacao': {'buffer_m': 50}})))
    assert list(com['camadas']['pontos']['ac_n_feicoes']) == [1, 0]


def test_geometria_invalida_e_camada_mista_sao_preparadas_e_nao_recusadas():
    gravata = Polygon([(X, Y), (X + 10, Y + 10), (X + 10, Y), (X, Y + 10), (X, Y)])
    mista = gdf([gravata, Point(X + 5, Y + 50), LineString([(X, Y + 60), (X + 10, Y + 60)])], nome=['p', 'q', 'r'])
    partes, estat = preparar(mista, 'mista')
    assert set(partes) == {0, 1, 2} and estat['corrigidas'] == 1 and estat['colapsadas'] == 0
    assert all(g.is_valid for g in partes[2].geometry)
    with pytest.raises(ValueError, match='separação por tipo'):
        preparar(mista, 'mista', separar=False)
    entrada = gdf([Point(X + 2, Y + 5)], proj_id=['P'])
    saida = enriquecer(entrada, tema(('t', mista, {'prefixo': 'tb'})))['camadas']['pontos']
    # Base mista vira uma parte por tipo, cada uma com prefixo próprio.
    assert {'tb_pol_nome', 'tb_pto_nome', 'tb_lin_nome'} <= set(saida.columns)
    assert saida.iloc[0]['tb_pol_n_feicoes'] == 1


def test_entrada_com_tipos_mistos_gera_uma_camada_por_tipo_e_crs_de_saida():
    entrada = gdf([Point(X, Y), LineString([(X, Y), (X + 10, Y)]), box(X, Y, X + 5, Y + 5)], proj_id=['a', 'b', 'c'])
    base = gdf([box(X - 1, Y - 1, X + 20, Y + 20)], nome=['B'])
    saida = enriquecer(entrada, tema(('b', base, {'prefixo': 'b'})))
    assert set(saida['camadas']) == {'pontos', 'linhas', 'poligonos'}
    assert all(str(c.crs) == 'EPSG:4674' for c in saida['camadas'].values())
    assert all(c.iloc[0]['b_nome'] == 'B' for c in saida['camadas'].values())


def test_campo_inexistente_e_prefixo_repetido_sao_recusados():
    entrada = gdf([Point(X, Y)], proj_id=['a'])
    base = gdf([box(X - 1, Y - 1, X + 1, Y + 1)], nome=['B'])
    with pytest.raises(ValueError, match='inexistente'):
        enriquecer(entrada, tema(('b', base, {'campos': ['nao_existe']})))
    with pytest.raises(ValueError, match='mais de uma base'):
        enriquecer(entrada, tema(('b1', base, {'prefixo': 'x'}), ('b2', base, {'prefixo': 'x'})))


def test_pacote_leva_camadas_apelidos_dicionario_e_configuracao(tmp_path):
    import io, json, zipfile
    import pyogrio
    from openpyxl import load_workbook
    from osgeo import ogr
    from api.services.extracao_atributos_pacote_enriquecimento import montar_pacote
    entrada = gdf([Point(X + 5, Y + 5), LineString([(X + 10, Y + 50), (X + 190, Y + 50)])],
                  proj_id=['P1', 'P2'], obs=['texto com \x0b controle', None])
    saida = enriquecer(entrada, tema(('municipios', MUNICIPIOS, {'papel': 'recorte', 'prefixo': 'mun',
                                                                 'apelidos': {'NM_MUN': 'Nome do município'}})))
    pacote, nome, manifesto = montar_pacote(saida['camadas'], entrada, saida['dicionario'],
                                            {'modo': 'enriquecimento'}, 'Teste enriquecimento')
    assert nome == 'teste_enriquecimento.zip'
    chaves = [item['chave'] for item in manifesto]
    assert chaves == ['gpkg', 'csv_pontos', 'csv_linhas', 'xlsx', 'csv_dicionario', 'configuracao', 'validacao']
    with zipfile.ZipFile(io.BytesIO(pacote)) as arquivo:
        (tmp_path / 'r.gpkg').write_bytes(arquivo.read(manifesto[0]['nome']))
        livro = load_workbook(io.BytesIO(arquivo.read(manifesto[3]['nome'])))
        dicionario_csv = arquivo.read(manifesto[4]['nome']).decode('utf-8-sig')
        configuracao = json.loads(arquivo.read(manifesto[5]['nome']))
    assert {n for n, _ in pyogrio.list_layers(tmp_path / 'r.gpkg')} == {'pontos', 'linhas', 'entrada'}
    assert pyogrio.read_info(tmp_path / 'r.gpkg', layer='linhas')['features'] == 2
    fonte = ogr.Open(str(tmp_path / 'r.gpkg'))
    definicao = fonte.GetLayerByName('linhas').GetLayerDefn()
    assert definicao.GetFieldDefn(definicao.GetFieldIndex('mun_NM_MUN')).GetAlternativeName() == 'Nome do município'
    fonte = None
    assert livro.sheetnames == ['pontos', 'linhas', 'dicionario_campos']
    assert 'mun_NM_MUN;Nome do município;Tema;municipios;NM_MUN;unidade de recorte' in dicionario_csv
    assert configuracao == {'modo': 'enriquecimento'}


@pytest.mark.parametrize('modo', ['enriquecimento', 'estatisticas'])
def test_execucao_do_servico_grava_camadas_pacote_e_finaliza_sem_erro(monkeypatch, modo):
    """Caminho completo do serviço com banco e gravação de camada simulados."""
    from contextlib import contextmanager
    from api.services import ciclo_vida_arquivos as ciclo
    from api.services import extracao_atributos as service
    from api.services import municipal_layer
    from api.services.extracao_atributos_regras import validar_conjunto
    entrada = gdf([LineString([(X + 10, Y + 50), (X + 190, Y + 50)]), Point(X + 5, Y + 5)], proj_id=['P1', 'P2'])
    frames = {'entrada': entrada, 'mun': MUNICIPIOS, 'uc': gdf([box(X, Y, X + 50, Y + 100)], nome=['UC'])}
    monkeypatch.setattr(municipal_layer, 'carregar_para_extracao', lambda ident: frames[ident])
    gravadas, inseridos, finalizacoes, usos = [], [], [], []
    monkeypatch.setattr(service.geo, 'registrar_camada',
                        lambda frame, nome, origem, **kw: gravadas.append((nome, len(frame))) or f'camada_{len(gravadas)}')

    class Conexao:
        def execute(self, sql, parametros=None):
            inseridos.append((sql, parametros))

    @contextmanager
    def conexao():
        yield Conexao()
    monkeypatch.setattr(service, 'get_connection', conexao)
    monkeypatch.setattr(ciclo, 'registrar_uso', lambda conn, camada, uso, ref: usos.append(camada))
    monkeypatch.setattr(ciclo, 'finalizar', lambda ident, **kw: finalizacoes.append(kw))
    categorias = validar_conjunto([
        {'id': 'adm', 'nome': 'Administrativo', 'conceito': '', 'camadas': [
            {'id': 'mun', 'nome': 'Municípios', 'regra': {'papel': 'recorte', 'prefixo': 'mun'}}]},
        {'id': 'amb', 'nome': 'Ambiental', 'conceito': '', 'camadas': [{'id': 'uc', 'nome': 'UC', 'regra': None}]}])
    params = {'camada_id': 'entrada', 'input_nome': 'Projetos', 'operacao': modo, 'categorias': categorias,
              'responsavel': 'teste', 'nome_saida': 'Projetos enriquecidos', 'opcoes': {}}
    service._execute('00000000-0000-0000-0000-000000000001', params)
    assert finalizacoes == [{}], finalizacoes
    assert sorted(nome for nome, _ in gravadas) == ['Projetos enriquecidos — linhas', 'Projetos enriquecidos — pontos']
    assert sorted(usos) == ['camada_1', 'camada_2']
    sql, valores = inseridos[-1]
    assert 'INSERT INTO geoprocessamento.extracao_atributos' in sql and valores[2] == modo
    relatorio = valores[8].obj
    assert relatorio['operacao'] == modo
    if modo == 'estatisticas':
        assert sum(n for _, n in gravadas) == len(entrada)
    assert relatorio['modo'] == 'enriquecimento' and set(relatorio['camadas']) == {'linhas', 'pontos'}
    assert relatorio['resumo']['camadas_intersectadas'] >= 1 and valores[11].endswith('.zip')


def test_varias_entradas_com_identificador_filtro_e_campos():
    pontos = gdf([Point(X + 5, Y + 5), Point(X + 150, Y + 5), Point(X + 50, Y + 50)],
                 proj_id=['A1', None, ' '], titulo=['a', 'b', 'c'], extra=[1, 2, 3])
    linhas = gdf([LineString([(X + 10, Y + 50), (X + 190, Y + 50)])], proj_id=['L1'], titulo=['linha'])
    entradas = [
        {'nome': 'Pontos1', 'frame': pontos,
         'config': {'campo_id': 'proj_id', 'filtro': {'campo': 'proj_id', 'operador': 'preenchido'},
                    'campos': ['titulo']}},
        {'nome': 'Linhas1', 'frame': linhas, 'config': {'campo_id': 'proj_id'}},
    ]
    saida = enriquecer(categorias=tema(('municipios', MUNICIPIOS, {'papel': 'recorte', 'prefixo': 'mun'})),
                       entradas=entradas)
    pontos_saida = saida['camadas']['pontos']
    # Só o ponto com proj_id preenchido; posição original e identificador preservados; 'extra' descartado.
    assert list(pontos_saida['camada_origem']) == ['Pontos1'] and list(pontos_saida['fid_origem']) == [0]
    assert list(pontos_saida['id_origem']) == ['A1'] and 'extra' not in pontos_saida.columns
    linhas_saida = saida['camadas']['linhas']
    assert set(linhas_saida['camada_origem']) == {'Linhas1'} and set(linhas_saida['id_origem']) == {'L1'}
    entradas_rel = {e['nome']: e for e in saida['relatorio']['entradas']}
    assert entradas_rel['Pontos1']['feicoes_selecionadas'] == 1 and entradas_rel['Pontos1']['feicoes_origem'] == 3


def test_validacao_aprova_resultado_correto_e_reprova_escolha_adulterada():
    from api.services import extracao_atributos_enriquecimento as motor
    entrada = gdf([LineString([(X, Y + 50), (X + 250, Y + 50)])], proj_id=['P1'])
    base = gdf([box(X - 10, Y + 40, X + 20, Y + 60), box(X + 20, Y + 40, X + 110, Y + 60)], nome=['pequena', 'grande'])
    categorias = tema(('municipios', MUNICIPIOS, {'papel': 'recorte', 'prefixo': 'mun'}),
                      ('uc', base, {'prefixo': 'uc'}))
    validacao = enriquecer(entrada, categorias)['relatorio']['validacao']
    linhas = validacao['linhas']
    assert validacao['aprovada'] and linhas['aprovada']
    assert linhas['feicoes_de_entrada_sem_registro'] == 0 and linhas['trechos_fora_da_propria_unidade'] == 0
    assert linhas['feicoes_com_soma_de_trechos_divergente'] == 0
    assert linhas['multiplicidade_conferida'][0]['divergencias'] == 0

    # Uma escolha trocada depois do enriquecimento precisa ser apanhada pela conferência.
    original = motor.enriquecer_base

    def adulterado(*args, **kwargs):
        tabela, info = original(*args, **kwargs)
        tabela[info['coluna_fid']] = tabela[info['coluna_fid']].map(lambda v: 0 if v == 1 else v)
        return tabela, info
    motor.enriquecer_base = adulterado
    try:
        reprovada = enriquecer(entrada, categorias)['relatorio']['validacao']
    finally:
        motor.enriquecer_base = original
    assert not reprovada['aprovada'] and reprovada['linhas']['multiplicidade_conferida'][0]['divergencias'] > 0


def test_finalidade_gera_subconjunto_de_campos_e_recusa_campo_inexistente():
    entrada = gdf([LineString([(X + 10, Y + 50), (X + 190, Y + 50)])], proj_id=['P1'], titulo=['t'])
    categorias = tema(('municipios', MUNICIPIOS, {'papel': 'recorte', 'prefixo': 'mun'}))
    saida = enriquecer(entrada, categorias, finalidades=[{'nome': 'Indicadores Dani', 'campos': ['mun_NM_MUN']}])
    finalidade = saida['finalidades']['indicadores_dani']
    assert list(finalidade['camadas']['linhas'].columns[:-1]) == [
        'id_registro', 'camada_origem', 'fid_origem', 'id_origem', 'mun_NM_MUN']
    with pytest.raises(ValueError, match='inexistente'):
        enriquecer(entrada, categorias, finalidades=[{'nome': 'X', 'campos': ['nao_existe']}])


def test_servico_recusa_varias_entradas_fora_do_enriquecimento(monkeypatch):
    from types import SimpleNamespace
    from api.services import extracao_atributos as service
    monkeypatch.setattr(service, 'catalogo', lambda: {
        'categorias': [{'id': 't', 'nome': 'T', 'conceito': ''}],
        'camadas': [{'id': 'a', 'nome': 'A'}, {'id': 'b', 'nome': 'B'}, {'id': 'c', 'nome': 'C'}]})
    usuario = SimpleNamespace(id='u')
    with pytest.raises(ValueError, match='só existem no modo enriquecimento'):
        service.iniciar({'input_id': 'a', 'operacao': 'intersection', 'categorias': [{'id': 't', 'camadas': ['c']}],
                         'entradas': [{'id': 'a'}, {'id': 'b'}]}, usuario)
    with pytest.raises(ValueError, match='igual à entrada'):
        service.iniciar({'input_id': 'a', 'operacao': 'enriquecimento', 'categorias': [{'id': 't', 'camadas': ['b']}],
                         'entradas': [{'id': 'a'}, {'id': 'b'}]}, usuario)


def test_campo_da_entrada_com_nome_reservado_nao_duplica_coluna():
    entrada = gdf([Point(X, Y)], fid_origem=['campo original'], id_registro=['outro'])
    base = gdf([box(X - 1, Y - 1, X + 1, Y + 1)], nome=['B'])
    saida = enriquecer(entrada, tema(('b', base, {'prefixo': 'b'})))['camadas']['pontos']
    assert saida.columns.is_unique
    assert saida.iloc[0]['fid_origem'] == 0 and saida.iloc[0]['fid_origem_2'] == 'campo original'
