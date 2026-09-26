"""Contrato sem recorte: cálculos independentes, geometria, pacote e API."""
import io
import json
import zipfile

import geopandas as gpd
import pandas as pd
import pytest
from shapely.geometry import Point, LineString, MultiLineString, Polygon, GeometryCollection, box

from api.services.extracao_atributos_estatisticas import enriquecer, agregar
from api.services.extracao_atributos_regras import normalizar_estatisticas

X, Y = 5_000_000, 7_500_000


def frame(geoms, **campos):
    return gpd.GeoDataFrame(campos, geometry=geoms, crs=5880)


def grupo(base, regra=None, id='social'):
    return [{'id': id, 'nome': id.title(), 'camadas': [{'id': 'b', 'nome': 'Base', 'frame': base,
                                                     'regra': {'prefixo': 'b', **(regra or {})}}]}]


@pytest.mark.parametrize('medida,esperado', [('media', 2), ('moda', 2), ('mediana', 2), ('total', 8),
    ('minimo', 0), ('maximo', 4), ('desvio_padrao', 2 ** .5), ('variancia', 2), ('contagem', 4)])
def test_nove_estatisticas_com_zero_e_nulos(medida, esperado):
    assert agregar([0, 2, 2, 4, None, float('nan')], medida) == pytest.approx(esperado)


def test_textos_empates_e_nulos():
    assert agregar(['B', 'A', 'A', 'B'], 'moda') == 'B'
    assert agregar(['0007', '0007', None], 'moda') == '0007'
    with pytest.raises(ValueError, match='não numérico'):
        agregar(['0007', '3'], 'media')
    assert agregar([None, pd.NA], 'media') is None
    assert agregar([None, pd.NA], 'contagem') == 0
    assert agregar([7], 'desvio_padrao') == 0
    assert agregar([7], 'variancia') == 0


def test_preserva_geometria_contagem_campos_e_intersecoes_independentes():
    entrada = frame([MultiLineString([[(X, Y+2), (X+25, Y+2)], [(X, Y+5), (X+25, Y+5)]]),
                     LineString([(X+100, Y), (X+110, Y)])], identificador=['dentro', 'fora'])
    base = frame([box(X, Y, X+10, Y+10), box(X+10, Y, X+30, Y+10)], numero=[0., 4.], texto=['A', 'B'])
    categorias = grupo(base, {'estatisticas_campos': {'numero':'media', 'texto': 'moda'}})
    categorias += [{'id': 'risco', 'nome': 'Risco', 'camadas': [{'id': 'r', 'nome': 'Risco', 'frame': base,
                     'regra': {'prefixo': 'r'}}]}]
    res = enriquecer(entrada, categorias)
    saida = res['camadas']['linhas']
    assert len(saida) == 2
    assert saida.geometry.to_wkb().equals(entrada.to_crs(4674).geometry.to_wkb())
    assert list(saida.identificador) == ['dentro', 'fora']
    assert list(saida.b_n_feicoes) == [2, 0], 'multipartes não contam a mesma base duas vezes'
    assert saida.iloc[0].b_numero == 2 and pd.isna(saida.iloc[1].b_numero)
    assert saida.iloc[0].b_texto == 'A' and pd.isna(saida.iloc[1].b_texto)
    assert json.loads(saida.iloc[0].r_numero) == [0.,4.] and json.loads(saida.iloc[0].r_texto) == ['A','B']
    assert pd.isna(saida.iloc[1].r_numero) and pd.isna(saida.iloc[1].r_texto)
    assert list(saida.r_intersecao) == ['Sim','Não']
    assert res['relatorio']['validacao']['geometrias_preservadas']


@pytest.mark.parametrize('categoria', ['risco', 'restricao'])
def test_binario_independe_de_atributo_nulo_e_toque_na_borda(categoria):
    entrada = frame([Point(X, Y), Point(X+100, Y)], codigo=['a', 'b'])
    base = frame([box(X, Y, X+10, Y+10)], vazio=[None])
    res = enriquecer(entrada, grupo(base, id=categoria))['camadas']['pontos']
    assert res.iloc[0].b_vazio == '[null]' and pd.isna(res.iloc[1].b_vazio)
    assert list(res.b_n_contato_borda) == [1,0]
    assert list(res.b_intersecao) == ['Sim', 'Não']


def test_base_vazia_mantem_campos_e_preserva_geometrias_nulas_e_invalidas():
    entrada = frame([None, Point(), GeometryCollection([Point(X, Y), LineString([(X, Y), (X+1, Y)])]),
                     Polygon([(X,Y), (X+10,Y+10), (X+10,Y), (X,Y+10), (X,Y)])], codigo=[1,2,3,4])
    base = frame([], valor=[])
    res = enriquecer(entrada, grupo(base))
    saida = pd.concat(res['camadas'].values()).sort_values('fid_origem')
    assert len(saida) == 4 and saida.b_valor.isna().all()
    assert list(saida.geometry.to_wkb()) == list(entrada.to_crs(4674).geometry.to_wkb())


def test_contagem_distingue_sem_intersecao_de_valores_nulos():
    entrada = frame([Point(X, Y), Point(X+100,Y)])
    base = frame([box(X-1,Y-1,X+1,Y+1)], v=[None])
    saida = enriquecer(entrada, grupo(base, {'estatisticas_campos': {'v':'contagem'}}))['camadas']['pontos']
    assert saida.iloc[0].b_v == 0 and pd.isna(saida.iloc[1].b_v)


def test_regra_antiga_nao_recorta_nem_aplica_buffer_no_novo_modo():
    entrada = frame([LineString([(X, Y), (X+20, Y)])])
    base = frame([box(X,Y,X+10,Y+10), box(X+10,Y,X+20,Y+10)], v=[1,3])
    res = enriquecer(entrada, grupo(base, {'papel':'recorte', 'multiplicidade':'todas',
                                         'preparacao': {'buffer_m':20}}))['camadas']['linhas']
    assert len(res) == 1 and res.iloc[0].b_v == '[1,3]'
    assert res.geometry.to_wkb().equals(entrada.to_crs(4674).geometry.to_wkb())


def test_campos_invalidos_e_filtros_nao_sao_silenciosos():
    entrada = frame([Point(X,Y)], id=['a'])
    base = frame([Point(X,Y)], v=[1])
    with pytest.raises(ValueError, match='inexistente'):
        enriquecer(entrada, grupo(base, {'estatisticas_campos': {'fantasma':'media'}}))
    with pytest.raises(ValueError, match='Remova o filtro'):
        enriquecer(categorias=grupo(base), entradas=[{'nome':'A','frame':entrada,
                   'config':{'filtro':{'campo':'id','operador':'preenchido'}}}])


def test_pacote_reaberto_tem_valores_e_geometrias(tmp_path):
    import pyogrio
    from openpyxl import load_workbook
    from api.services.extracao_atributos_pacote_enriquecimento import montar_pacote
    entrada = frame([Point(X,Y), Point(X+100,Y)], codigo=['a','b'])
    base = frame([box(X-1,Y-1,X+1,Y+1), box(X-2,Y-2,X+2,Y+2)], valor=[2,6])
    resultado = enriquecer(entrada, grupo(base, {'estatisticas_campos': {'valor':'total'}}))
    pacote, _, arquivos = montar_pacote(resultado['camadas'], entrada, resultado['dicionario'],
                                         {'operacao':'estatisticas'}, 'Teste', preservar_geometrias=True, validacao=resultado['relatorio']['validacao'])
    with zipfile.ZipFile(io.BytesIO(pacote)) as z:
        gpkg = tmp_path/'s.gpkg';gpkg.write_bytes(z.read(arquivos[0]['nome']))
        livro = load_workbook(io.BytesIO(z.read(next(a['nome'] for a in arquivos if a['chave']=='xlsx'))))
    saida = pyogrio.read_dataframe(gpkg, layer='pontos')
    assert [r['atributos']['valor'] for r in json.loads(saida.iloc[0].b_correspondencias)] == [2,6]
    assert len(saida)==2 and saida.iloc[0].b_valor==8 and pd.isna(saida.iloc[1].b_valor)
    assert saida.geometry.to_wkb().equals(entrada.to_crs(4674).geometry.to_wkb())
    cab = [c.value for c in livro['pontos'][1]]
    assert livro['pontos'].cell(2, cab.index('b_valor')+1).value==8
    assert livro['pontos'].cell(3, cab.index('b_valor')+1).value is None


def test_api_admite_os_dois_fluxos_e_rejeita_medida_desconhecida():
    from pydantic import ValidationError
    from api.routers.extracao_atributos import Extracao
    for modo in ['enriquecimento','estatisticas']:
        pedido = Extracao(input_id='a', operacao=modo, categorias=[{'id':'social','camadas':['b'],
                           'regras':{'b':{'estatistica':'moda','estatisticas_campos':{'v':'variancia'}}}}])
        assert pedido.operacao == modo
    with pytest.raises(ValidationError):
        Extracao(input_id='a',operacao='estatisticas',categorias=[{'id':'social','camadas':['b'],
                 'regras':{'b':{'estatistica':'inexistente'}}}])


def test_persistencia_wkb_preserva_vazias_e_precisao():
    import shapely
    from api.repositories.camada_geoespacial_repository import _feature_rows, _insert_features
    entrada = frame([Point(), None, Point(X+.123456789, Y+.987654321)])
    rows = _feature_rows(entrada, usar_wkb=True)
    assert [None if wkb is None else shapely.from_wkb(wkb).wkb for _, _, wkb in rows] == list(entrada.to_crs(4674).geometry.to_wkb())
    calls = []
    class Cursor:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def executemany(self, sql, rows): calls.append((sql.as_string(), rows))
    class Conn:
        def cursor(self): return Cursor()
    _insert_features(Conn(), 'camada_processada_feicao', 'id', rows, usar_wkb=True)
    assert 'ST_GeomFromWKB' in calls[0][0]
    assert len(calls[0][1]) == 3
    assert calls[0][1][0][3] is not None and calls[0][1][1][3] is None


def test_configuracao_salva_medidas_por_campo_e_modo(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from api.services import configuracao_bancada as config, extracao_atributos as service
    monkeypatch.setattr(config, 'raiz', lambda: tmp_path)
    monkeypatch.setattr(service, 'catalogo', lambda: {'categorias':[{'id':'social','nome':'Social'}],
                       'camadas':[{'id':'e','nome':'Entrada'},{'id':'b','nome':'Base'}]})
    config.salvar('Estatísticas', [{'id':'social','camadas':['b'], 'regras':{'b':{'estatistica':'mediana',
                  'estatisticas_campos':{'nome':'moda','valor':'total'}}}}], SimpleNamespace(id='teste'),
                  entradas=[{'id':'e'}], operacao='estatisticas')
    salvo = config.carregar('estatisticas')
    assert salvo['operacao'] == 'estatisticas'
    regra = salvo['categorias'][0]['camadas'][0]['regra']
    assert regra['estatistica']=='mediana' and regra['estatisticas_campos']=={'nome':'moda','valor':'total'}
